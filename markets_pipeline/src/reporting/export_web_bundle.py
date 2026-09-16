from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from src.common.dollar_linked import DOLLAR_LINKED_SYMBOLS
from src.common.monetary_variables import (
    MONETARY_CORE_METRICS,
    MONETARY_RATES_TNA_LABELS,
    MONETARY_VARIABLE_LABELS,
)


ROOT = Path(__file__).resolve().parents[2]
PROCESSED_DIR = ROOT / "data" / "processed"
WEB_DIR = ROOT / "web"
WEB_DATA_DIR = WEB_DIR / "data"
OUTPUT_FILE = WEB_DATA_DIR / "dashboard_payload.json"
OUTPUT_JS_FILE = WEB_DATA_DIR / "dashboard_payload.js"

LATEST_FILE = PROCESSED_DIR / "latest_snapshot.csv"
SERIES_FILE = PROCESSED_DIR / "market_series_daily.csv"
LECAP_FILE = PROCESSED_DIR / "lecap_monitor_latest.csv"
LECAP_HISTORY_FILE = PROCESSED_DIR / "lecap_monitor_history.csv"
CER_CURVE_FILE = PROCESSED_DIR / "cer_curve_latest.csv"
CER_CURVE_HISTORY_FILE = PROCESSED_DIR / "cer_curve_history.csv"
HARD_DOLLAR_CURVE_FILE = PROCESSED_DIR / "hard_dollar_curve_latest.csv"
HARD_DOLLAR_CURVE_HISTORY_FILE = PROCESSED_DIR / "hard_dollar_curve_history.csv"
HARD_DOLLAR_WEIGHTED_PARITY_HISTORY_FILE = PROCESSED_DIR / "hard_dollar_weighted_parity_history.csv"
FX_FUTURES_CURVE_LATEST_FILE = PROCESSED_DIR / "fx_futures_curve_latest.csv"
DOLLAR_LINKED_HISTORY_FILE = PROCESSED_DIR / "dollar_linked_volume_history.csv"
DOLLAR_LINKED_LATEST_FILE = PROCESSED_DIR / "dollar_linked_volume_latest.csv"
MONETARY_VARIABLES_LATEST_FILE = PROCESSED_DIR / "monetary_variables_latest.csv"
MONETARY_VARIABLES_ECONOMY_PIB_HISTORY_FILE = PROCESSED_DIR / "monetary_variables_economy_pib_history.csv"
MONETARY_VARIABLES_PESO_USD_PIB_HISTORY_FILE = PROCESSED_DIR / "monetary_variables_peso_usd_pib_history.csv"
MONETARY_VARIABLES_USD_STOCKS_HISTORY_FILE = PROCESSED_DIR / "monetary_variables_usd_stocks_history.csv"
MONETARY_VARIABLES_BLOCK_FRESHNESS_FILE = PROCESSED_DIR / "monetary_variables_block_freshness.csv"
MONETARY_VARIABLES_SOURCE_TABLE_FILE = PROCESSED_DIR / "monetary_variables_source_table.csv"
MONETARY_RATES_TNA_SERIES_FILE = PROCESSED_DIR / "monetary_rates_tna_series.csv"
MONETARY_RATES_TNA_REFERENCE_FILE = PROCESSED_DIR / "monetary_rates_tna_reference.csv"
MONETARY_RATES_TNA_LATEST_FILE = PROCESSED_DIR / "monetary_rates_tna_latest.csv"
FCI_CATEGORY_HISTORY_FILE = PROCESSED_DIR / "fci_category_patrimony_history.csv"
FCI_CURRENCY_HISTORY_FILE = PROCESSED_DIR / "fci_currency_patrimony_history.csv"
FCI_HORIZON_HISTORY_FILE = PROCESSED_DIR / "fci_horizon_patrimony_history.csv"
FCI_MONEY_MARKET_LATEST_FILE = PROCESSED_DIR / "fci_money_market_latest.csv"
FCI_REPORTED_SLICE_HISTORY_FILE = PROCESSED_DIR / "fci_reported_slice_history.csv"
BREAKEVEN_FILE = PROCESSED_DIR / "cer_breakeven_monthly_latest.csv"
LECAP_MACRO_HISTORY_FILE = PROCESSED_DIR / "lecap_macro_history.csv"
LECAP_MACRO_MONTHLY_FILE = PROCESSED_DIR / "lecap_macro_monthly.csv"
ALERTS_FILE = PROCESSED_DIR / "alerts_latest.csv"
META_FILE = PROCESSED_DIR / "pipeline_metadata.json"

HERO_ASSETS = [
    "usd_ars_ccl",
    "merval_usd",
    "country_risk",
    "lecap_3m_yield",
]
EXTERNAL_ASSET_ORDER = [
    "sp500",
    "nasdaq",
    "russell2000",
    "nikkei225",
    "hang_seng",
    "us_2y",
    "us_10y",
    "usd_broad",
    "eurusd",
    "brent",
    "gold",
    "copper",
    "bitcoin",
]
EXTERNAL_PANEL_DEFS = [
    ("equities", "Acciones", ["sp500", "nasdaq", "russell2000", "nikkei225", "hang_seng"]),
    ("rates", "Tasas", ["us_2y", "us_10y"]),
    ("fx", "FX global", ["usd_broad", "eurusd"]),
    ("commodities_crypto", "Commodities y cripto", ["brent", "gold", "copper", "bitcoin"]),
]
MARKET_ASSET_SERIES = [
    "merval_usd",
    "usd_ars_ccl",
    "usd_ars_official",
    "country_risk",
]
HARD_DOLLAR_A = ["AO27D", "AO28D", "AL29D", "AN29D", "AL30D", "AL35D", "AE38D", "AL41D"]
HARD_DOLLAR_G = ["GD29D", "GD30D", "GD35D", "GD38D", "GD41D", "GD46D"]
SEVERITY_ORDER = {"critical": 0, "warning": 1, "info": 2}
CATEGORY_ORDER = ["FX", "Rates", "Equities", "Sovereign Risk"]
BCRA_BAND_URL = "https://www.bcra.gob.ar/regimen-de-bandas-cambiarias/"
BCRA_BAND_XLSX_URL = "https://www.bcra.gob.ar/archivos/Pdfs/PublicacionesEstadisticas/serie-completa-bandas-cambiarias.xlsx"
ECOGO_INFLATION_PROJECTION_FILE = ROOT / "legacy" / "ecogo_projection_inputs" / "proy_infla_ecogo.xlsx"
SPREADS_START_DATE = pd.Timestamp("2025-05-01")
REGIME_START_DATE = pd.Timestamp("2025-04-14")
REGIME_START_UPPER = 1400.0
REGIME_2025_MONTHLY_UPPER_RATE = 0.01
FCI_MATRIX_WINDOWS: list[tuple[str, str, pd.Timedelta]] = [
    ("change_1d", "1D", pd.Timedelta(days=1)),
    ("change_7d", "7D", pd.Timedelta(days=7)),
    ("change_30d", "30D", pd.Timedelta(days=30)),
    ("change_90d", "90D", pd.Timedelta(days=90)),
    ("change_ytd", "YTD", None),
    ("change_365d", "365D", pd.Timedelta(days=365)),
]
FCI_REPORTED_DIMENSION_ORDER = ["benchmark", "duration", "currency", "region", "mixed_type", "fund_type"]
FCI_REPORTED_CATEGORY_ORDER = ["renta_fija", "mercado_dinero", "renta_mixta", "renta_variable"]


def read_csv(path: Path) -> pd.DataFrame:
    if not path.exists():
        return pd.DataFrame()
    return pd.read_csv(path)


def float_or_none(value: object) -> float | None:
    if pd.isna(value):
        return None
    return float(value)


def format_pct(value: object) -> float | None:
    if pd.isna(value):
        return None
    return round(float(value) * 100.0, 3)


def safe_str(value: object) -> str | None:
    if pd.isna(value):
        return None
    return str(value)


def display_source_name(source: object, source_detail: object = None) -> str | None:
    primary = safe_str(source)
    if primary and primary not in {"fred_external_prices", "yahoo_external_prices"}:
        return primary
    detail = safe_str(source_detail) or ""
    if detail.startswith("fred_graph:"):
        return "FRED"
    if detail.startswith("yahoo_chart:"):
        return "Yahoo Finance"
    return primary


def build_freshness(
    latest: pd.DataFrame,
    lecap_history: pd.DataFrame,
    cer_curve_history: pd.DataFrame,
    hard_dollar_history: pd.DataFrame,
    fx_futures_curve_latest: pd.DataFrame,
    dollar_linked_history: pd.DataFrame,
    monetary_variables_latest: pd.DataFrame,
    fci_category_history: pd.DataFrame,
    fci_currency_history: pd.DataFrame,
    fci_horizon_history: pd.DataFrame,
    fci_money_market_latest: pd.DataFrame,
    fci_reported_slice_history: pd.DataFrame,
) -> dict[str, object]:
    def latest_date(frame: pd.DataFrame, column: str) -> str | None:
        if frame.empty or column not in frame.columns:
            return None
        parsed = pd.to_datetime(frame[column], errors="coerce").dropna()
        if parsed.empty:
            return None
        return parsed.max().strftime("%Y-%m-%d")

    local_latest = latest[latest["category"].isin(CATEGORY_ORDER)].copy()
    external_latest = latest[latest["category"] == "External"].copy()
    blocks = [
        {"block": "overview", "label": "Portada", "latest_date": latest_date(local_latest, "date")},
        {"block": "external_context", "label": "Contexto externo", "latest_date": latest_date(external_latest, "date")},
        {"block": "fixed_curve", "label": "Curva del Tesoro", "latest_date": latest_date(lecap_history, "date")},
        {"block": "cer_curve", "label": "Curva CER", "latest_date": latest_date(cer_curve_history, "date")},
        {"block": "hard_dollar", "label": "Hard-dollar", "latest_date": latest_date(hard_dollar_history, "date")},
        {"block": "fx_futures", "label": "Futuros DLR", "latest_date": latest_date(fx_futures_curve_latest, "as_of_date")},
        {"block": "dollar_linked", "label": "Dolar linked", "latest_date": latest_date(dollar_linked_history, "date")},
        {"block": "monetary_variables", "label": "Monetarias", "latest_date": latest_date(monetary_variables_latest, "date")},
        {"block": "fci_category", "label": "FCI por categoría", "latest_date": latest_date(fci_category_history, "date")},
        {"block": "fci_currency", "label": "FCI por moneda", "latest_date": latest_date(fci_currency_history, "date")},
        {"block": "fci_horizon", "label": "FCI por horizonte", "latest_date": latest_date(fci_horizon_history, "date")},
        {"block": "fci_money_market", "label": "Money market leaders", "latest_date": latest_date(fci_money_market_latest, "date")},
        {"block": "fci_reported_slices", "label": "FCI cortes reportados", "latest_date": latest_date(fci_reported_slice_history, "date")},
    ]
    latest_dates = [item["latest_date"] for item in blocks if item["latest_date"]]
    return {
        "reference_date": max(latest_dates) if latest_dates else None,
        "blocks": blocks,
    }


def build_overview(latest: pd.DataFrame) -> dict[str, object]:
    latest = latest[latest["category"].isin(CATEGORY_ORDER)].copy()
    latest["ret_1d_pct"] = latest["ret_1d"].map(format_pct)
    latest["ret_1m_pct"] = latest["ret_1m"].map(format_pct)
    latest["ret_ytd_pct"] = latest["ret_ytd"].map(format_pct)
    latest["zscore"] = latest["zscore_252"].map(float_or_none)
    latest["value"] = latest["value_nominal"].map(float_or_none)
    latest["date"] = latest["date"].astype(str)

    categories: list[dict[str, object]] = []
    for category in CATEGORY_ORDER:
        category_df = latest[latest["category"] == category].sort_values("asset_label").copy()
        if category_df.empty:
            continue
        categories.append(
            {
                "category": category,
                "assets": [
                    {
                        "asset": row.asset,
                        "label": row.asset_label,
                        "unit": row.unit,
                        "value": float_or_none(row.value),
                        "date": row.date,
                        "performance_basis": safe_str(getattr(row, "performance_basis", None)),
                        "ret_1d_pct": float_or_none(row.ret_1d_pct),
                        "ret_1m_pct": float_or_none(row.ret_1m_pct),
                        "ret_ytd_pct": float_or_none(row.ret_ytd_pct),
                        "zscore": float_or_none(row.zscore),
                    }
                    for row in category_df.itertuples(index=False)
                ],
            }
        )
    return {
        "latest_date": safe_str(latest["date"].max()) if not latest.empty else None,
        "asset_count": int(len(latest)),
        "category_count": int(latest["category"].nunique()) if not latest.empty else 0,
        "categories": categories,
    }


def build_hero_metrics(latest: pd.DataFrame) -> list[dict[str, object]]:
    hero_df = latest[latest["asset"].isin(HERO_ASSETS)].copy()
    hero_df["ret_1d_pct"] = hero_df["ret_1d"].map(format_pct)
    hero_df["ret_ytd_pct"] = hero_df["ret_ytd"].map(format_pct)
    hero_df["value"] = hero_df["value_nominal"].map(float_or_none)
    hero_df = hero_df.sort_values("asset_label")
    return [
        {
            "asset": row.asset,
            "label": row.asset_label,
            "value": float_or_none(row.value),
            "unit": row.unit,
            "performance_basis": safe_str(getattr(row, "performance_basis", None)),
            "ret_1d_pct": float_or_none(row.ret_1d_pct),
            "ret_ytd_pct": float_or_none(row.ret_ytd_pct),
        }
        for row in hero_df.itertuples(index=False)
    ]


def build_series(history: pd.DataFrame) -> dict[str, list[dict[str, object]]]:
    history = history.copy()
    history["date"] = pd.to_datetime(history["date"], errors="coerce")
    history = history[history["asset"].isin(HERO_ASSETS)].copy()
    cutoff = history["date"].max() - pd.Timedelta(days=180)
    history = history[history["date"] >= cutoff].copy()
    series_payload: dict[str, list[dict[str, object]]] = {}
    for asset, asset_df in history.groupby("asset"):
        asset_df = asset_df.sort_values("date")
        series_payload[asset] = [
            {
                "date": row.date.strftime("%Y-%m-%d"),
                "index": float_or_none(row.index_base_100),
                "value": float_or_none(row.value_nominal),
            }
            for row in asset_df.itertuples(index=False)
        ]
    return series_payload


def build_market_asset_series(history: pd.DataFrame, days: int = 420) -> dict[str, list[dict[str, object]]]:
    history = history.copy()
    history["date"] = pd.to_datetime(history["date"], errors="coerce")
    history["value_nominal"] = pd.to_numeric(history["value_nominal"], errors="coerce")
    history = history[history["asset"].isin(MARKET_ASSET_SERIES)].dropna(subset=["date", "value_nominal"]).copy()
    if history.empty:
        return {}

    cutoff = history["date"].max() - pd.Timedelta(days=days)
    history = history[history["date"] >= cutoff].copy()

    payload: dict[str, list[dict[str, object]]] = {}
    for asset, asset_df in history.groupby("asset", sort=True):
        asset_df = asset_df.sort_values("date")
        payload[asset] = [
            {
                "date": row.date.strftime("%Y-%m-%d"),
                "value": float_or_none(row.value_nominal),
            }
            for row in asset_df.itertuples(index=False)
        ]
    return payload


def build_external_context(latest: pd.DataFrame, history: pd.DataFrame, days: int = 180) -> dict[str, object]:
    latest = latest[latest["category"] == "External"].copy()
    history = history[history["category"] == "External"].copy()
    latest["ret_1d_pct"] = latest["ret_1d"].map(format_pct)
    latest["ret_1m_pct"] = latest["ret_1m"].map(format_pct)
    latest["ret_ytd_pct"] = latest["ret_ytd"].map(format_pct)
    latest["value"] = latest["value_nominal"].map(float_or_none)
    latest["date"] = pd.to_datetime(latest["date"], errors="coerce")
    history["date"] = pd.to_datetime(history["date"], errors="coerce")
    history["value_nominal"] = pd.to_numeric(history["value_nominal"], errors="coerce")
    history["index_base_100"] = pd.to_numeric(history["index_base_100"], errors="coerce")

    latest_lookup = {row.asset: row for row in latest.itertuples(index=False)}
    if history["date"].notna().any():
        cutoff = history["date"].dropna().max() - pd.Timedelta(days=days)
        history = history[history["date"] >= cutoff].copy()

    series_payload: dict[str, list[dict[str, object]]] = {}
    for asset, asset_df in history.groupby("asset", sort=True):
        asset_df = asset_df.sort_values("date")
        series_payload[asset] = [
            {
                "date": row.date.strftime("%Y-%m-%d"),
                "value": float_or_none(row.value_nominal),
                "index": float_or_none(row.index_base_100),
            }
            for row in asset_df.itertuples(index=False)
            if pd.notna(row.date)
        ]

    def serialize_asset(asset_key: str) -> dict[str, object] | None:
        row = latest_lookup.get(asset_key)
        if row is None:
            return None
        return {
            "asset": row.asset,
            "label": row.asset_label,
            "unit": row.unit,
            "value": float_or_none(row.value),
            "date": row.date.strftime("%Y-%m-%d") if pd.notna(row.date) else None,
            "ret_1d_pct": float_or_none(row.ret_1d_pct),
            "ret_1m_pct": float_or_none(row.ret_1m_pct),
            "ret_ytd_pct": float_or_none(row.ret_ytd_pct),
            "source": display_source_name(getattr(row, "source", None), getattr(row, "source_detail", None)),
            "source_url": safe_str(getattr(row, "source_url", None)),
        }

    benchmark_strip = [item for item in (serialize_asset(asset) for asset in EXTERNAL_ASSET_ORDER) if item]
    panels = [
        {
            "panel": panel_key,
            "label": panel_label,
            "items": [item for item in (serialize_asset(asset) for asset in panel_assets) if item],
        }
        for panel_key, panel_label, panel_assets in EXTERNAL_PANEL_DEFS
    ]

    latest_dates = [item["date"] for item in benchmark_strip if item.get("date")]
    return {
        "latest_date": max(latest_dates) if latest_dates else None,
        "asset_count": len(benchmark_strip),
        "benchmark_strip": benchmark_strip,
        "panels": panels,
        "series": series_payload,
    }


def build_fixed_curve(lecap_latest: pd.DataFrame) -> list[dict[str, object]]:
    frame = lecap_latest.copy()
    return [
        {
            "symbol": row.symbol,
            "family": row.instrument_family,
            "tir_pct": float_or_none(row.tea_pct),
            "duration": float_or_none(row.modified_duration),
            "maturity_date": safe_str(row.maturity_date),
            "payment_date": safe_str(row.payment_date),
            "volume": float_or_none(row.volume),
        }
        for row in frame.itertuples(index=False)
    ]


def build_fixed_curve_history(lecap_history: pd.DataFrame, max_dates: int = 12) -> dict[str, object]:
    if lecap_history.empty:
        return {"dates": [], "curves": {}}

    frame = lecap_history.copy()
    frame["date"] = pd.to_datetime(frame["date"], errors="coerce")
    frame["duration"] = pd.to_numeric(frame["modified_duration"], errors="coerce")
    frame["tir_pct"] = pd.to_numeric(frame["tea_pct"], errors="coerce")
    frame["volume"] = pd.to_numeric(frame["volume"], errors="coerce")
    frame = frame.dropna(subset=["date", "duration", "tir_pct"]).sort_values(["date", "duration", "symbol"])
    selected_dates = frame["date"].dt.strftime("%Y-%m-%d").drop_duplicates().tolist()[-max_dates:]
    curves: dict[str, list[dict[str, object]]] = {}
    for date_key in selected_dates:
        day_df = frame[frame["date"].dt.strftime("%Y-%m-%d") == date_key].copy()
        curves[date_key] = [
            {
                "symbol": row.symbol,
                "family": row.instrument_family,
                "tir_pct": float_or_none(row.tir_pct),
                "duration": float_or_none(row.duration),
                "maturity_date": safe_str(row.maturity_date),
                "payment_date": safe_str(row.payment_date),
                "volume": float_or_none(row.volume),
            }
            for row in day_df.itertuples(index=False)
        ]
    return {"dates": selected_dates, "curves": curves}


def build_cer_curve(cer_curve_latest: pd.DataFrame) -> list[dict[str, object]]:
    frame = cer_curve_latest.copy()
    return [
        {
            "symbol": row.symbol,
            "tir_pct": format_pct(row.tir),
            "duration": float_or_none(row.duration),
            "payment_date": safe_str(row.payment_date),
            "price": float_or_none(row.price),
            "technical_price": float_or_none(row.technical_price),
            "parity": float_or_none(row.parity),
            "volume": float_or_none(row.volume),
        }
        for row in frame.itertuples(index=False)
    ]


def build_cer_curve_history(cer_curve_history: pd.DataFrame, max_dates: int = 12) -> dict[str, object]:
    if cer_curve_history.empty:
        return {"dates": [], "curves": {}}

    frame = cer_curve_history.copy()
    frame["date"] = pd.to_datetime(frame["date"], errors="coerce")
    frame["duration"] = pd.to_numeric(frame["duration"], errors="coerce")
    frame["tir"] = pd.to_numeric(frame["tir"], errors="coerce")
    frame["volume"] = pd.to_numeric(frame["volume"], errors="coerce")
    frame["price"] = pd.to_numeric(frame["price"], errors="coerce")
    frame["parity"] = pd.to_numeric(frame["parity"], errors="coerce")
    frame["technical_price"] = pd.to_numeric(frame["technical_price"], errors="coerce")
    frame = frame.dropna(subset=["date", "duration", "tir"]).sort_values(["date", "duration", "symbol"])
    selected_dates = frame["date"].dt.strftime("%Y-%m-%d").drop_duplicates().tolist()[-max_dates:]
    curves: dict[str, list[dict[str, object]]] = {}
    for date_key in selected_dates:
        day_df = frame[frame["date"].dt.strftime("%Y-%m-%d") == date_key].copy()
        curves[date_key] = [
            {
                "symbol": row.symbol,
                "tir_pct": format_pct(row.tir),
                "duration": float_or_none(row.duration),
                "payment_date": safe_str(row.payment_date),
                "price": float_or_none(row.price),
                "technical_price": float_or_none(row.technical_price),
                "parity": float_or_none(row.parity),
                "volume": float_or_none(row.volume),
            }
            for row in day_df.itertuples(index=False)
        ]
    return {"dates": selected_dates, "curves": curves}


def latest_observed_inflation_month(observed_monthly: pd.DataFrame | None) -> pd.Timestamp | None:
    if observed_monthly is None or observed_monthly.empty or "month" not in observed_monthly.columns:
        return None

    frame = observed_monthly.copy()
    frame["month"] = pd.to_datetime(frame["month"], errors="coerce").dt.to_period("M").dt.to_timestamp()
    value_column = None
    for candidate in ("inflation_monthly_pct", "inflation_monthly", "value_pct"):
        if candidate in frame.columns:
            value_column = candidate
            break
    if value_column is not None:
        frame[value_column] = pd.to_numeric(frame[value_column], errors="coerce")
        frame = frame.dropna(subset=["month", value_column])
    else:
        frame = frame.dropna(subset=["month"])
    if frame.empty:
        return None
    return pd.Timestamp(frame["month"].max()).normalize()


def build_breakeven_curve(
    breakeven: pd.DataFrame,
    observed_monthly: pd.DataFrame | None = None,
) -> list[dict[str, object]]:
    frame = breakeven.sort_values("month").copy()
    frame["month"] = pd.to_datetime(frame["month"], errors="coerce")
    observed_cutoff = latest_observed_inflation_month(observed_monthly)
    if observed_cutoff is not None:
        frame = frame[frame["month"] > observed_cutoff].copy()
    return [
        {
            "month": row.month.strftime("%Y-%m-%d") if not pd.isna(row.month) else None,
            "plot_date": ((row.month + pd.offsets.MonthEnd(0)).strftime("%Y-%m-%d")) if not pd.isna(row.month) else None,
            "month_label": row.month_label,
            "pair": row.source_pair_label,
            "value_pct": float_or_none(row.forward_monthly_pct),
            "gap_fill": bool(row.is_gap_fill),
            "is_latent": bool(getattr(row, "is_latent", False)),
        }
        for row in frame.itertuples(index=False)
    ]


def build_alerts(alerts: pd.DataFrame) -> list[dict[str, object]]:
    frame = alerts.copy()
    frame["severity_rank"] = frame["severity"].map(SEVERITY_ORDER).fillna(99)
    frame = frame.sort_values(["severity_rank", "category", "asset_label"]).head(8)
    return [
        {
            "date": safe_str(row.as_of_date),
            "category": row.category,
            "asset": row.asset,
            "asset_label": row.asset_label,
            "severity": row.severity,
            "metric_key": safe_str(row.metric_key),
            "metric_label": safe_str(row.metric_label),
            "observed_value": float_or_none(row.observed_value),
            "threshold_value": float_or_none(row.threshold_value),
            "direction": safe_str(row.direction),
        }
        for row in frame.itertuples(index=False)
    ]


def build_hard_dollar_curve(hard_dollar_latest: pd.DataFrame) -> list[dict[str, object]]:
    frame = hard_dollar_latest.copy()
    if frame.empty:
        return []
    return [
        {
            "symbol": row.symbol,
            "family": row.family,
            "tir_pct": format_pct(row.tir),
            "duration": float_or_none(row.modified_duration),
            "payment_date": safe_str(row.payment_date),
            "price": float_or_none(row.price),
            "technical_price": float_or_none(row.technical_price),
            "parity": float_or_none(row.parity),
            "volume": float_or_none(row.volume),
        }
        for row in frame.itertuples(index=False)
    ]


def build_hard_dollar_curve_history(hard_dollar_history: pd.DataFrame, max_dates: int = 12) -> dict[str, object]:
    if hard_dollar_history.empty:
        return {"dates": [], "curves": {}}

    frame = hard_dollar_history.copy()
    frame["date"] = pd.to_datetime(frame["date"], errors="coerce")
    frame["duration"] = pd.to_numeric(frame["modified_duration"], errors="coerce")
    frame["tir"] = pd.to_numeric(frame["tir"], errors="coerce")
    frame["volume"] = pd.to_numeric(frame["volume"], errors="coerce")
    frame["price"] = pd.to_numeric(frame["price"], errors="coerce")
    frame["parity"] = pd.to_numeric(frame["parity"], errors="coerce")
    frame["technical_price"] = pd.to_numeric(frame["technical_price"], errors="coerce")
    frame = frame.dropna(subset=["date", "duration", "tir"]).sort_values(["date", "family", "duration", "symbol"])
    selected_dates = frame["date"].dt.strftime("%Y-%m-%d").drop_duplicates().tolist()[-max_dates:]
    curves: dict[str, list[dict[str, object]]] = {}
    for date_key in selected_dates:
        day_df = frame[frame["date"].dt.strftime("%Y-%m-%d") == date_key].copy()
        curves[date_key] = [
            {
                "symbol": row.symbol,
                "family": row.family,
                "tir_pct": format_pct(row.tir),
                "duration": float_or_none(row.duration),
                "payment_date": safe_str(row.payment_date),
                "price": float_or_none(row.price),
                "technical_price": float_or_none(row.technical_price),
                "parity": float_or_none(row.parity),
                "volume": float_or_none(row.volume),
            }
            for row in day_df.itertuples(index=False)
        ]
    return {"dates": selected_dates, "curves": curves}


def build_hard_dollar_weighted_parity(weighted_history: pd.DataFrame) -> dict[str, list[dict[str, object]]]:
    if weighted_history.empty:
        return {"AL": [], "GD": []}

    frame = weighted_history.copy()
    frame["date"] = pd.to_datetime(frame["date"], errors="coerce")
    frame["weighted_parity"] = pd.to_numeric(frame["weighted_parity"], errors="coerce")
    frame["total_outstanding"] = pd.to_numeric(frame["total_outstanding"], errors="coerce")
    frame["bond_count"] = pd.to_numeric(frame["bond_count"], errors="coerce")
    frame = frame.dropna(subset=["date", "family", "weighted_parity"]).sort_values(["date", "family"])

    payload: dict[str, list[dict[str, object]]] = {"AL": [], "GD": []}
    for family in payload.keys():
        family_frame = frame[frame["family"].astype(str).str.upper() == family].copy()
        payload[family] = [
            {
                "date": row.date.strftime("%Y-%m-%d"),
                "weighted_parity": float_or_none(row.weighted_parity),
                "total_outstanding": float_or_none(row.total_outstanding),
                "bond_count": int(row.bond_count) if not pd.isna(row.bond_count) else None,
            }
            for row in family_frame.itertuples(index=False)
        ]
    return payload


def build_hard_dollar_country_risk(history: pd.DataFrame) -> list[dict[str, object]]:
    if history.empty:
        return []
    frame = history.copy()
    frame["date"] = pd.to_datetime(frame["date"], errors="coerce")
    frame["value_nominal"] = pd.to_numeric(frame["value_nominal"], errors="coerce")
    frame = frame[frame["asset"] == "country_risk"].dropna(subset=["date", "value_nominal"]).copy()
    frame = frame.sort_values("date")
    return [
        {
            "date": row.date.strftime("%Y-%m-%d"),
            "value": float_or_none(row.value_nominal),
        }
        for row in frame.itertuples(index=False)
    ]


def build_fx_futures_curve(fx_futures_curve_latest: pd.DataFrame) -> list[dict[str, object]]:
    if fx_futures_curve_latest.empty:
        return []
    frame = fx_futures_curve_latest.copy()
    frame["as_of_date"] = pd.to_datetime(frame["as_of_date"], errors="coerce")
    frame["maturity_date"] = pd.to_datetime(frame["maturity_date"], errors="coerce")
    for column in ("settlement", "volume", "open_interest", "implied_rate"):
        frame[column] = pd.to_numeric(frame[column], errors="coerce")
    frame = frame.dropna(subset=["as_of_date", "maturity_date", "symbol", "settlement"]).copy()
    frame = frame[frame["maturity_date"] >= pd.Timestamp.today().normalize()].copy()
    if frame.empty:
        return []
    frame = frame.sort_values(["maturity_date", "symbol"]).reset_index(drop=True)
    return [
        {
            "as_of_date": row.as_of_date.strftime("%Y-%m-%d"),
            "symbol": row.symbol,
            "contract_month": safe_str(row.contract_month),
            "maturity_date": row.maturity_date.strftime("%Y-%m-%d"),
            "settlement": float_or_none(row.settlement),
            "volume": float_or_none(row.volume),
            "open_interest": float_or_none(row.open_interest),
            "implied_rate": float_or_none(row.implied_rate),
            "source_url": safe_str(row.source_url),
        }
        for row in frame.itertuples(index=False)
    ]


def build_dollar_linked_monitor(
    dollar_linked_history: pd.DataFrame,
    dollar_linked_latest: pd.DataFrame,
) -> dict[str, object]:
    history_frame = dollar_linked_history.copy()
    latest_frame = dollar_linked_latest.copy()

    if "date" in history_frame.columns:
        history_frame["date"] = pd.to_datetime(history_frame["date"], errors="coerce")
    if "price" in history_frame.columns:
        history_frame["price"] = pd.to_numeric(history_frame["price"], errors="coerce")
    if "volume" in history_frame.columns:
        history_frame["volume"] = pd.to_numeric(history_frame["volume"], errors="coerce")
    history_frame["symbol"] = history_frame.get("symbol", pd.Series(dtype="object")).astype(str).str.upper()
    history_frame = history_frame[history_frame["symbol"].isin(DOLLAR_LINKED_SYMBOLS)].copy()
    history_frame["symbol_rank"] = history_frame["symbol"].map({symbol: idx for idx, symbol in enumerate(DOLLAR_LINKED_SYMBOLS)})
    history_frame = history_frame.sort_values(["date", "symbol_rank"]).reset_index(drop=True)

    if "date" in latest_frame.columns:
        latest_frame["date"] = pd.to_datetime(latest_frame["date"], errors="coerce")
    latest_frame["symbol"] = latest_frame.get("symbol", pd.Series(dtype="object")).astype(str).str.upper()
    latest_frame["price"] = pd.to_numeric(latest_frame.get("price"), errors="coerce")
    latest_frame["volume"] = pd.to_numeric(latest_frame.get("volume"), errors="coerce")
    latest_frame["amount"] = pd.to_numeric(latest_frame.get("amount"), errors="coerce")
    latest_frame = latest_frame[latest_frame["symbol"].isin(DOLLAR_LINKED_SYMBOLS)].copy()
    latest_frame["symbol_rank"] = latest_frame["symbol"].map({symbol: idx for idx, symbol in enumerate(DOLLAR_LINKED_SYMBOLS)})
    latest_frame = latest_frame.sort_values("symbol_rank").reset_index(drop=True)

    history_rows = [
        {
            "date": row.date.strftime("%Y-%m-%d") if not pd.isna(row.date) else None,
            "symbol": row.symbol,
            "price": float_or_none(row.price),
            "volume": float_or_none(row.volume),
        }
        for row in history_frame.itertuples(index=False)
    ]
    latest_rows = [
        {
            "date": row.date.strftime("%Y-%m-%d") if not pd.isna(row.date) else None,
            "symbol": row.symbol,
            "price": float_or_none(row.price),
            "volume": float_or_none(row.volume),
            "amount": float_or_none(row.amount),
        }
        for row in latest_frame.itertuples(index=False)
    ]

    latest_dates = [
        *pd.to_datetime(history_frame.get("date", pd.Series(dtype="datetime64[ns]")), errors="coerce").dropna().tolist(),
        *pd.to_datetime(latest_frame.get("date", pd.Series(dtype="datetime64[ns]")), errors="coerce").dropna().tolist(),
    ]
    latest_date = max(latest_dates).strftime("%Y-%m-%d") if latest_dates else None

    return {
        "latest_date": latest_date,
        "symbols": list(DOLLAR_LINKED_SYMBOLS),
        "history": history_rows,
        "latest": latest_rows,
    }


def build_monetary_variables(
    monetary_latest: pd.DataFrame,
    economy_pib_history: pd.DataFrame,
    peso_usd_pib_history: pd.DataFrame,
    usd_stocks_history: pd.DataFrame,
    block_freshness: pd.DataFrame,
    source_table: pd.DataFrame,
    rates_tna_series: pd.DataFrame,
    rates_tna_reference: pd.DataFrame,
    rates_tna_latest: pd.DataFrame,
) -> dict[str, object]:
    latest_frame = monetary_latest.copy()
    if "date" in latest_frame.columns:
        latest_frame["date"] = pd.to_datetime(latest_frame["date"], errors="coerce")
    latest_frame["metric"] = latest_frame.get("metric", pd.Series(dtype="object")).astype(str)
    label_fallback = latest_frame["label"] if "label" in latest_frame.columns else latest_frame["metric"]
    latest_frame["label"] = latest_frame["metric"].map(MONETARY_VARIABLE_LABELS).fillna(label_fallback)
    latest_frame["value"] = pd.to_numeric(latest_frame.get("value"), errors="coerce")
    latest_frame["change_30d_pct"] = pd.to_numeric(latest_frame.get("change_30d_pct"), errors="coerce")
    latest_frame["change_yoy_pct"] = pd.to_numeric(latest_frame.get("change_yoy_pct"), errors="coerce")
    latest_frame["lag_days"] = pd.to_numeric(latest_frame.get("lag_days"), errors="coerce")
    latest_frame["metric_rank"] = latest_frame["metric"].map({metric: idx for idx, metric in enumerate(MONETARY_CORE_METRICS)})
    latest_frame = latest_frame[latest_frame["metric"].isin(MONETARY_CORE_METRICS)].sort_values("metric_rank").reset_index(drop=True)

    latest_dates = [
        *pd.to_datetime(latest_frame.get("date", pd.Series(dtype="datetime64[ns]")), errors="coerce").dropna().tolist(),
        *pd.to_datetime(economy_pib_history.get("month", pd.Series(dtype="datetime64[ns]")), errors="coerce").dropna().tolist(),
        *pd.to_datetime(peso_usd_pib_history.get("month", pd.Series(dtype="datetime64[ns]")), errors="coerce").dropna().tolist(),
        *pd.to_datetime(usd_stocks_history.get("month", pd.Series(dtype="datetime64[ns]")), errors="coerce").dropna().tolist(),
        *pd.to_datetime(rates_tna_latest.get("date", pd.Series(dtype="datetime64[ns]")), errors="coerce").dropna().tolist(),
    ]
    latest_date = max(latest_dates).strftime("%Y-%m-%d") if latest_dates else None

    def serialize_monthly(frame: pd.DataFrame) -> list[dict[str, object]]:
        rows: list[dict[str, object]] = []
        if frame.empty:
            return rows
        local = frame.copy()
        local["month"] = pd.to_datetime(local.get("month"), errors="coerce")
        for row in local.itertuples(index=False):
            payload_row: dict[str, object] = {}
            for key, value in row._asdict().items():
                if isinstance(value, pd.Timestamp):
                    payload_row[key] = value.strftime("%Y-%m-%d") if not pd.isna(value) else None
                elif pd.isna(value):
                    payload_row[key] = None
                elif isinstance(value, (int, float)):
                    payload_row[key] = float(value)
                else:
                    payload_row[key] = value
            rows.append(payload_row)
        return rows

    freshness_frame = block_freshness.copy()
    if "latest_date" in freshness_frame.columns:
        freshness_frame["latest_date"] = pd.to_datetime(freshness_frame["latest_date"], errors="coerce")
    freshness_rows = [
        {
            "block": row.block,
            "label": row.label,
            "latest_date": row.latest_date.strftime("%Y-%m-%d") if not pd.isna(row.latest_date) else None,
            "window_days": float_or_none(row.window_days),
            "lag_days": float_or_none(row.lag_days),
        }
        for row in freshness_frame.itertuples(index=False)
    ]

    source_frame = source_table.copy()
    if "date" in source_frame.columns:
        source_frame["date"] = pd.to_datetime(source_frame["date"], errors="coerce")
    else:
        source_frame["date"] = pd.Series(dtype="datetime64[ns]")
    if "label" not in source_frame.columns:
        source_frame["label"] = source_frame.get("series_key", pd.Series(dtype="object")).astype(str)
    source_rows = [
        {
            "series_key": row.series_key,
            "label": row.label,
            "date": row.date.strftime("%Y-%m-%d") if not pd.isna(row.date) else None,
            "value": float_or_none(row.value),
            "unit": safe_str(row.unit),
            "lag_days": float_or_none(row.lag_days),
        }
        for row in source_frame.sort_values(["date", "label"]).itertuples(index=False)
    ]

    rates_series_frame = rates_tna_series.copy()
    if "date" in rates_series_frame.columns:
        rates_series_frame["date"] = pd.to_datetime(rates_series_frame["date"], errors="coerce")
    else:
        rates_series_frame["date"] = pd.Series(dtype="datetime64[ns]")
    rates_series_frame["metric"] = rates_series_frame.get("metric", pd.Series(dtype="object")).astype(str)
    rates_series_label_fallback = rates_series_frame["label"] if "label" in rates_series_frame.columns else rates_series_frame["metric"]
    rates_series_frame["label"] = rates_series_frame["metric"].map(MONETARY_RATES_TNA_LABELS).fillna(rates_series_label_fallback)
    rates_series_frame["value"] = pd.to_numeric(rates_series_frame.get("value"), errors="coerce")
    rates_series_rows = [
        {
            "date": row.date.strftime("%Y-%m-%d") if not pd.isna(row.date) else None,
            "metric": row.metric,
            "label": row.label,
            "value": float_or_none(row.value),
        }
        for row in rates_series_frame.sort_values(["metric", "date"]).itertuples(index=False)
    ]

    rates_reference_frame = rates_tna_reference.copy()
    if "date" in rates_reference_frame.columns:
        rates_reference_frame["date"] = pd.to_datetime(rates_reference_frame["date"], errors="coerce")
    else:
        rates_reference_frame["date"] = pd.Series(dtype="datetime64[ns]")
    if "metric" not in rates_reference_frame.columns:
        rates_reference_frame["metric"] = pd.Series(dtype="object")
    if "label" not in rates_reference_frame.columns:
        rates_reference_frame["label"] = rates_reference_frame["metric"].astype(str)
    rates_reference_frame["value"] = pd.to_numeric(rates_reference_frame.get("value"), errors="coerce")
    rates_reference_rows = [
        {
            "date": row.date.strftime("%Y-%m-%d") if not pd.isna(row.date) else None,
            "metric": row.metric,
            "label": row.label,
            "value": float_or_none(row.value),
            "source_type": safe_str(row.source_type),
        }
        for row in rates_reference_frame.sort_values(["metric", "date"]).itertuples(index=False)
    ]

    rates_latest_frame = rates_tna_latest.copy()
    if "date" in rates_latest_frame.columns:
        rates_latest_frame["date"] = pd.to_datetime(rates_latest_frame["date"], errors="coerce")
    else:
        rates_latest_frame["date"] = pd.Series(dtype="datetime64[ns]")
    rates_latest_frame["metric"] = rates_latest_frame.get("metric", pd.Series(dtype="object")).astype(str)
    rates_latest_label_fallback = rates_latest_frame["label"] if "label" in rates_latest_frame.columns else rates_latest_frame["metric"]
    rates_latest_frame["label"] = rates_latest_frame["metric"].map(MONETARY_RATES_TNA_LABELS).fillna(rates_latest_label_fallback)
    rates_latest_frame["value"] = pd.to_numeric(rates_latest_frame.get("value"), errors="coerce")
    rates_latest_frame["lag_days"] = pd.to_numeric(rates_latest_frame.get("lag_days"), errors="coerce")
    rates_latest_frame["spread_vs_reference_pct"] = pd.to_numeric(rates_latest_frame.get("spread_vs_reference_pct"), errors="coerce")
    rates_latest_rows = [
        {
            "metric": row.metric,
            "label": row.label,
            "date": row.date.strftime("%Y-%m-%d") if not pd.isna(row.date) else None,
            "value": float_or_none(row.value),
            "lag_days": float_or_none(row.lag_days),
            "spread_vs_reference_pct": float_or_none(row.spread_vs_reference_pct),
        }
        for row in rates_latest_frame.sort_values(["date", "label"]).itertuples(index=False)
    ]
    rates_block_freshness = next((row for row in freshness_rows if row["block"] == "rates_tna"), None)

    return {
        "latest_date": latest_date,
        "core_latest": [
            {
                "metric": row.metric,
                "label": row.label,
                "date": row.date.strftime("%Y-%m-%d") if not pd.isna(row.date) else None,
                "value": float_or_none(row.value),
                "unit": safe_str(getattr(row, "unit", None)),
                "change_30d_pct": float_or_none(row.change_30d_pct),
                "change_yoy_pct": float_or_none(row.change_yoy_pct),
                "lag_days": float_or_none(row.lag_days),
            }
            for row in latest_frame.itertuples(index=False)
        ],
        "economy_pib_history": serialize_monthly(economy_pib_history),
        "peso_usd_pib_history": serialize_monthly(peso_usd_pib_history),
        "usd_stocks_history": serialize_monthly(usd_stocks_history),
        "rates_tna": {
            "latest_date": max(
                (
                    item["date"]
                    for item in rates_latest_rows
                    if item.get("date")
                ),
                default=None,
            ),
            "series": rates_series_rows,
            "reference_series": rates_reference_rows,
            "latest": rates_latest_rows,
            "block_freshness": rates_block_freshness,
        },
        "block_freshness": freshness_rows,
        "source_table": source_rows,
    }


def build_fci_category_history(fci_category_history: pd.DataFrame) -> dict[str, object]:
    if fci_category_history.empty:
        return {"series": [], "latest": []}

    frame = fci_category_history.copy()
    frame["date"] = pd.to_datetime(frame["date"], errors="coerce")
    frame["patrimony"] = pd.to_numeric(frame["patrimony"], errors="coerce")
    frame["fund_count"] = pd.to_numeric(frame["fund_count"], errors="coerce")
    frame = frame.dropna(subset=["date", "patrimony"]).sort_values(["date", "category_key"]).copy()
    if frame.empty:
        return {"series": [], "latest": []}

    latest_date = frame["date"].max()
    latest = frame[frame["date"] == latest_date].copy()
    prior_dates = frame.loc[frame["date"] < latest_date, "date"].drop_duplicates().sort_values()
    previous_date = prior_dates.iloc[-1] if not prior_dates.empty else pd.NaT
    previous = frame[frame["date"] == previous_date][["category_key", "patrimony"]].rename(columns={"patrimony": "previous_patrimony"}) if not pd.isna(previous_date) else pd.DataFrame(columns=["category_key", "previous_patrimony"])
    latest = latest.merge(previous, on="category_key", how="left")
    latest["change_pct"] = (latest["patrimony"] / latest["previous_patrimony"] - 1.0) * 100.0

    return {
        "series": [
            {
                "date": row.date.strftime("%Y-%m-%d"),
                "category_key": row.category_key,
                "category_label": row.category_label,
                "patrimony": float_or_none(row.patrimony),
                "fund_count": float_or_none(row.fund_count),
            }
            for row in frame.itertuples(index=False)
        ],
        "latest": [
            {
                "date": row.date.strftime("%Y-%m-%d"),
                "category_key": row.category_key,
                "category_label": row.category_label,
                "patrimony": float_or_none(row.patrimony),
                "fund_count": float_or_none(row.fund_count),
                "change_pct": float_or_none(row.change_pct),
            }
            for row in latest.sort_values("patrimony", ascending=False).itertuples(index=False)
        ],
    }


def build_fci_currency_history(fci_currency_history: pd.DataFrame) -> dict[str, object]:
    if fci_currency_history.empty:
        return {"series": [], "latest": []}

    frame = fci_currency_history.copy()
    frame["date"] = pd.to_datetime(frame["date"], errors="coerce")
    frame["patrimony"] = pd.to_numeric(frame["patrimony"], errors="coerce")
    frame["row_count"] = pd.to_numeric(frame["row_count"], errors="coerce")
    frame = frame.dropna(subset=["date", "patrimony"]).sort_values(["date", "currency_label"]).copy()
    if frame.empty:
        return {"series": [], "latest": []}

    latest_date = frame["date"].max()
    latest = frame[frame["date"] == latest_date].copy()
    previous_date = frame.loc[frame["date"] < latest_date, "date"].drop_duplicates().sort_values()
    previous_date = previous_date.iloc[-1] if not previous_date.empty else pd.NaT
    previous = (
        frame[frame["date"] == previous_date][["currency_label", "patrimony"]].rename(columns={"patrimony": "previous_patrimony"})
        if not pd.isna(previous_date)
        else pd.DataFrame(columns=["currency_label", "previous_patrimony"])
    )
    latest = latest.merge(previous, on="currency_label", how="left")
    latest["change_pct"] = (latest["patrimony"] / latest["previous_patrimony"] - 1.0) * 100.0

    return {
        "series": [
            {
                "date": row.date.strftime("%Y-%m-%d"),
                "currency_label": row.currency_label,
                "patrimony": float_or_none(row.patrimony),
                "row_count": float_or_none(row.row_count),
            }
            for row in frame.itertuples(index=False)
        ],
        "latest": [
            {
                "date": row.date.strftime("%Y-%m-%d"),
                "currency_label": row.currency_label,
                "patrimony": float_or_none(row.patrimony),
                "row_count": float_or_none(row.row_count),
                "change_pct": float_or_none(row.change_pct),
            }
            for row in latest.sort_values("patrimony", ascending=False).itertuples(index=False)
        ],
    }


def build_fci_horizon_history(fci_horizon_history: pd.DataFrame) -> dict[str, object]:
    if fci_horizon_history.empty:
        return {"series": [], "latest": []}

    frame = fci_horizon_history.copy()
    frame["date"] = pd.to_datetime(frame["date"], errors="coerce")
    frame["patrimony"] = pd.to_numeric(frame["patrimony"], errors="coerce")
    frame["fund_count"] = pd.to_numeric(frame["fund_count"], errors="coerce")
    frame = frame.dropna(subset=["date", "patrimony"]).sort_values(["date", "horizon"]).copy()
    if frame.empty:
        return {"series": [], "latest": []}

    latest_date = frame["date"].max()
    latest = frame[frame["date"] == latest_date].copy()
    previous_date = frame.loc[frame["date"] < latest_date, "date"].drop_duplicates().sort_values()
    previous_date = previous_date.iloc[-1] if not previous_date.empty else pd.NaT
    previous = (
        frame[frame["date"] == previous_date][["horizon", "patrimony"]].rename(columns={"patrimony": "previous_patrimony"})
        if not pd.isna(previous_date)
        else pd.DataFrame(columns=["horizon", "previous_patrimony"])
    )
    latest = latest.merge(previous, on="horizon", how="left")
    latest["change_pct"] = (latest["patrimony"] / latest["previous_patrimony"] - 1.0) * 100.0

    return {
        "series": [
            {
                "date": row.date.strftime("%Y-%m-%d"),
                "horizon": row.horizon,
                "patrimony": float_or_none(row.patrimony),
                "fund_count": float_or_none(row.fund_count),
            }
            for row in frame.itertuples(index=False)
        ],
        "latest": [
            {
                "date": row.date.strftime("%Y-%m-%d"),
                "horizon": row.horizon,
                "patrimony": float_or_none(row.patrimony),
                "fund_count": float_or_none(row.fund_count),
                "change_pct": float_or_none(row.change_pct),
            }
            for row in latest.sort_values("patrimony", ascending=False).itertuples(index=False)
        ],
    }


def build_fci_money_market_latest(fci_money_market_latest: pd.DataFrame) -> list[dict[str, object]]:
    if fci_money_market_latest.empty:
        return []
    frame = fci_money_market_latest.copy()
    frame["date"] = pd.to_datetime(frame["date"], errors="coerce")
    frame["patrimony"] = pd.to_numeric(frame["patrimony"], errors="coerce")
    frame["vcp"] = pd.to_numeric(frame["vcp"], errors="coerce")
    frame["ccp"] = pd.to_numeric(frame["ccp"], errors="coerce")
    frame = frame.dropna(subset=["date"]).copy()
    frame = frame.sort_values("patrimony", ascending=False).head(20)
    return [
        {
            "date": row.date.strftime("%Y-%m-%d"),
            "fund": row.fund,
            "patrimony": float_or_none(row.patrimony),
            "vcp": float_or_none(row.vcp),
            "ccp": float_or_none(row.ccp),
            "horizon": safe_str(row.horizon),
            "source_url": safe_str(row.source_url),
        }
        for row in frame.itertuples(index=False)
    ]


def build_fci_aum_matrix(
    fci_category_history: pd.DataFrame,
    fci_currency_history: pd.DataFrame,
    fci_horizon_history: pd.DataFrame,
) -> dict[str, object]:
    def normalize(frame: pd.DataFrame, label_col: str, count_col: str | None) -> pd.DataFrame:
        if frame.empty:
            columns = ["date", label_col, "patrimony"]
            if count_col:
                columns.append(count_col)
            return pd.DataFrame(columns=columns)
        working = frame.copy()
        working["date"] = pd.to_datetime(working["date"], errors="coerce")
        working["patrimony"] = pd.to_numeric(working["patrimony"], errors="coerce")
        if count_col:
            working[count_col] = pd.to_numeric(working[count_col], errors="coerce")
        subset = ["date", label_col, "patrimony"]
        return working.dropna(subset=subset).sort_values(["date", label_col]).reset_index(drop=True)

    def get_snapshot(frame: pd.DataFrame, label_col: str, target_date: pd.Timestamp) -> pd.DataFrame:
        eligible = frame[frame["date"] <= target_date].copy()
        if eligible.empty:
            return pd.DataFrame(columns=[label_col, "base_patrimony"])
        snapshot = (
            eligible.sort_values(["date", label_col])
            .groupby(label_col, as_index=False)
            .tail(1)
            [[label_col, "patrimony"]]
            .rename(columns={"patrimony": "base_patrimony"})
        )
        return snapshot

    category = normalize(fci_category_history, "category_label", "fund_count")
    currency = normalize(fci_currency_history, "currency_label", "row_count")
    horizon = normalize(fci_horizon_history, "horizon", "fund_count")
    if category.empty:
        return {"latest_date": None, "reference_total": None, "rows": []}

    latest_date = category["date"].max()
    latest_category = category[category["date"] == latest_date].copy()
    total_patrimony = float(latest_category["patrimony"].sum()) if not latest_category.empty else None
    total_funds = float(latest_category["fund_count"].fillna(0).sum()) if "fund_count" in latest_category.columns else None

    def build_rows(
        frame: pd.DataFrame,
        *,
        section_key: str,
        section_label: str,
        label_col: str,
        label_alias: str,
        count_col: str | None = None,
    ) -> list[dict[str, object]]:
        if frame.empty:
            return []
        latest = frame[frame["date"] == latest_date].copy()
        if latest.empty:
            return []
        latest = latest.sort_values("patrimony", ascending=False).copy()
        total_for_section = float(latest["patrimony"].sum()) if not latest.empty else None
        for change_key, _, offset in FCI_MATRIX_WINDOWS:
            if change_key == "change_ytd":
                target_date = pd.Timestamp(year=latest_date.year - 1, month=12, day=31)
            else:
                target_date = latest_date - offset
            base = get_snapshot(frame, label_col, target_date)
            latest = latest.merge(base, on=label_col, how="left")
            latest[change_key] = latest["patrimony"] - latest["base_patrimony"]
            latest = latest.drop(columns=["base_patrimony"])
        rows: list[dict[str, object]] = []
        for row in latest.itertuples(index=False):
            rows.append(
                {
                    "section_key": section_key,
                    "section_label": section_label,
                    "label": getattr(row, label_col),
                    "label_alias": label_alias,
                    "date": latest_date.strftime("%Y-%m-%d"),
                    "patrimony": float_or_none(row.patrimony),
                    "share_pct": ((float(row.patrimony) / total_patrimony) * 100.0) if total_patrimony else None,
                    "section_share_pct": ((float(row.patrimony) / total_for_section) * 100.0) if total_for_section else None,
                    "count": float_or_none(getattr(row, count_col)) if count_col else None,
                    **{change_key: float_or_none(getattr(row, change_key)) for change_key, _, _ in FCI_MATRIX_WINDOWS},
                }
            )
        return rows

    rows = [
        {
            "section_key": "overview",
            "section_label": "Universo FCI",
            "label": "Total",
            "label_alias": "Total",
            "date": latest_date.strftime("%Y-%m-%d"),
            "patrimony": total_patrimony,
            "share_pct": 100.0 if total_patrimony else None,
            "section_share_pct": 100.0 if total_patrimony else None,
            "count": total_funds,
            **{
                change_key: float_or_none(
                    total_patrimony - get_snapshot(category, "category_label", pd.Timestamp(year=latest_date.year - 1, month=12, day=31))["base_patrimony"].sum()
                ) if change_key == "change_ytd" else None
                for change_key, _, _ in FCI_MATRIX_WINDOWS
            },
        }
    ]

    for change_key, _, offset in FCI_MATRIX_WINDOWS:
        if change_key == "change_ytd":
            target_date = pd.Timestamp(year=latest_date.year - 1, month=12, day=31)
        else:
            target_date = latest_date - offset
        prior = get_snapshot(category, "category_label", target_date)
        rows[0][change_key] = float_or_none(total_patrimony - prior["base_patrimony"].sum()) if total_patrimony is not None and not prior.empty else None

    rows.extend(
        build_rows(
            category,
            section_key="category",
            section_label="Categoría",
            label_col="category_label",
            label_alias="Fondos",
            count_col="fund_count",
        )
    )
    rows.extend(
        build_rows(
            horizon,
            section_key="horizon",
            section_label="Horizonte",
            label_col="horizon",
            label_alias="Fondos",
            count_col="fund_count",
        )
    )
    rows.extend(
        build_rows(
            currency,
            section_key="currency",
            section_label="Moneda",
            label_col="currency_label",
            label_alias="Registros",
            count_col="row_count",
        )
    )
    return {
        "latest_date": latest_date.strftime("%Y-%m-%d"),
        "reference_total": total_patrimony,
        "rows": rows,
    }


def build_fci_reported_slices(
    fci_reported_slice_history: pd.DataFrame,
    fci_category_history: pd.DataFrame,
) -> dict[str, object]:
    def get_snapshot(frame: pd.DataFrame, keys: list[str], target_date: pd.Timestamp) -> pd.DataFrame:
        eligible = frame[frame["date"] <= target_date].copy()
        if eligible.empty:
            return pd.DataFrame(columns=[*keys, "base_patrimony"])
        snapshot = (
            eligible.sort_values(["date", *keys])
            .groupby(keys, as_index=False)
            .tail(1)
            [[*keys, "patrimony"]]
            .rename(columns={"patrimony": "base_patrimony"})
        )
        return snapshot

    if fci_reported_slice_history.empty or fci_category_history.empty:
        return {"latest_date": None, "latest": [], "available_dimensions": [], "available_categories": []}

    frame = fci_reported_slice_history.copy()
    frame["date"] = pd.to_datetime(frame["date"], errors="coerce")
    frame["patrimony"] = pd.to_numeric(frame["patrimony"], errors="coerce")
    frame = frame.dropna(subset=["date", "patrimony"]).copy()
    frame = frame[frame["patrimony"] > 0].sort_values(
        ["date", "category_key", "dimension_key", "patrimony", "slice_label"],
        ascending=[True, True, True, False, True],
    )
    if frame.empty:
        return {"latest_date": None, "latest": [], "available_dimensions": [], "available_categories": []}

    category = fci_category_history.copy()
    category["date"] = pd.to_datetime(category["date"], errors="coerce")
    category["patrimony"] = pd.to_numeric(category["patrimony"], errors="coerce")
    category = category.dropna(subset=["date", "patrimony"]).copy()

    latest_date = frame["date"].max()
    latest = frame[frame["date"] == latest_date].copy()
    if latest.empty:
        return {"latest_date": None, "latest": [], "available_dimensions": [], "available_categories": []}

    category_latest = (
        category[category["date"] == latest_date][["category_key", "patrimony"]]
        .rename(columns={"patrimony": "category_total"})
        .drop_duplicates(subset=["category_key"])
    )
    latest = latest.merge(category_latest, on="category_key", how="left")
    latest["dimension_total"] = latest.groupby(["category_key", "dimension_key"])["patrimony"].transform("sum")
    latest["category_share_pct"] = (latest["patrimony"] / latest["category_total"]) * 100.0
    latest["dimension_share_pct"] = (latest["patrimony"] / latest["dimension_total"]) * 100.0

    keys = ["category_key", "dimension_key", "slice_label"]
    for change_key, _, offset in FCI_MATRIX_WINDOWS:
        if change_key == "change_ytd":
            target_date = pd.Timestamp(year=latest_date.year - 1, month=12, day=31)
        else:
            target_date = latest_date - offset
        base = get_snapshot(frame, keys, target_date)
        latest = latest.merge(base, on=keys, how="left")
        latest[change_key] = latest["patrimony"] - latest["base_patrimony"]
        latest = latest.drop(columns=["base_patrimony"])

    latest["category_rank"] = latest["category_key"].map({key: idx for idx, key in enumerate(FCI_REPORTED_CATEGORY_ORDER)}).fillna(999)
    latest["dimension_rank"] = latest["dimension_key"].map({key: idx for idx, key in enumerate(FCI_REPORTED_DIMENSION_ORDER)}).fillna(999)
    latest = latest.sort_values(
        ["category_rank", "dimension_rank", "patrimony", "slice_label"],
        ascending=[True, True, False, True],
    ).reset_index(drop=True)

    return {
        "latest_date": latest_date.strftime("%Y-%m-%d"),
        "available_dimensions": latest["dimension_key"].dropna().astype(str).drop_duplicates().tolist(),
        "available_categories": latest["category_key"].dropna().astype(str).drop_duplicates().tolist(),
        "latest": [
            {
                "date": row.date.strftime("%Y-%m-%d"),
                "category_key": row.category_key,
                "category_label": row.category_label,
                "dimension_key": row.dimension_key,
                "dimension_label": row.dimension_label,
                "slice_label": row.slice_label,
                "patrimony": float_or_none(row.patrimony),
                "category_total": float_or_none(row.category_total),
                "dimension_total": float_or_none(row.dimension_total),
                "category_share_pct": float_or_none(row.category_share_pct),
                "dimension_share_pct": float_or_none(row.dimension_share_pct),
                **{change_key: float_or_none(getattr(row, change_key)) for change_key, _, _ in FCI_MATRIX_WINDOWS},
            }
            for row in latest.itertuples(index=False)
        ],
    }


def build_hard_dollar(latest: pd.DataFrame, hard_dollar_latest: pd.DataFrame) -> dict[str, list[dict[str, object]]]:
    reference = hard_dollar_latest.copy()
    if reference.empty:
        return {"a_curve": [], "g_curve": []}

    reference["ret_1d_pct"] = pd.NA
    if not latest.empty:
        sov = latest[latest["category"] == "Sovereign Risk"].copy()
        sov["symbol"] = sov["asset_label"]
        sov["ret_1d_pct"] = sov["ret_1d"].map(format_pct)
        reference = reference.drop(columns=["ret_1d_pct"], errors="ignore").merge(
            sov[["symbol", "ret_1d_pct"]],
            on="symbol",
            how="left",
        )

    def collect(symbols: list[str]) -> list[dict[str, object]]:
        rows = reference[reference["symbol"].isin(symbols)].copy()
        rows["curve_order"] = rows["symbol"].map({symbol: idx for idx, symbol in enumerate(symbols)}).fillna(999)
        rows = rows.sort_values(["curve_order", "symbol"])
        return [
            {
                "symbol": row.symbol,
                "price": float_or_none(row.price),
                "tir_pct": format_pct(row.tir),
                "duration": float_or_none(row.modified_duration),
                "technical_price": float_or_none(row.technical_price),
                "parity": float_or_none(row.parity),
                "payment_date": safe_str(row.payment_date),
                "ret_1d_pct": float_or_none(row.ret_1d_pct),
            }
            for row in rows.itertuples(index=False)
        ]

    return {"a_curve": collect(HARD_DOLLAR_A), "g_curve": collect(HARD_DOLLAR_G)}


def derive_2025_upper_band_series(reference_dates: pd.Series | pd.Index | list[object]) -> pd.DataFrame:
    if reference_dates is None:
        return pd.DataFrame(columns=["date", "band_upper", "source_type"])

    dates = pd.to_datetime(pd.Series(reference_dates), errors="coerce").dropna().sort_values().unique()
    dates = pd.DatetimeIndex(dates)
    dates = dates[(dates >= SPREADS_START_DATE) & (dates < pd.Timestamp("2026-01-01"))]
    if dates.empty:
        return pd.DataFrame(columns=["date", "band_upper", "source_type"])

    max_date = dates.max()
    daily_values: dict[pd.Timestamp, float] = {REGIME_START_DATE: REGIME_START_UPPER}
    current_value = REGIME_START_UPPER
    for current_date in pd.date_range(REGIME_START_DATE + pd.Timedelta(days=1), max_date, freq="D"):
        month_factor = (1.0 + REGIME_2025_MONTHLY_UPPER_RATE) ** (1.0 / current_date.days_in_month)
        current_value *= month_factor
        daily_values[pd.Timestamp(current_date)] = current_value

    return pd.DataFrame(
        {
            "date": list(dates),
            "band_upper": [daily_values.get(pd.Timestamp(date)) for date in dates],
            "source_type": "Regla oficial 2025",
        }
    ).dropna(subset=["band_upper"])


def read_ecogo_inflation_projection() -> pd.DataFrame:
    if not ECOGO_INFLATION_PROJECTION_FILE.exists():
        return pd.DataFrame(columns=["month", "inflation_monthly"])

    workbook = load_workbook(ECOGO_INFLATION_PROJECTION_FILE, data_only=True, read_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    workbook.close()

    frame = pd.DataFrame(rows, columns=["month", "inflation_monthly"])
    frame["month"] = pd.to_datetime(frame["month"], errors="coerce")
    frame["inflation_monthly"] = pd.to_numeric(frame["inflation_monthly"], errors="coerce")
    frame = frame.dropna(subset=["month", "inflation_monthly"]).copy()
    return frame.sort_values("month").reset_index(drop=True)


def project_upper_band_with_ecogo(
    start_date: pd.Timestamp,
    start_value: float,
    end_date: pd.Timestamp,
    inflation_projection: pd.DataFrame,
) -> pd.DataFrame:
    if pd.isna(start_date) or pd.isna(end_date) or end_date <= start_date or pd.isna(start_value):
        return pd.DataFrame(columns=["date", "band_upper", "source_type"])
    if inflation_projection.empty:
        return pd.DataFrame(columns=["date", "band_upper", "source_type"])

    projection = inflation_projection.copy()
    projection["month_start"] = projection["month"].dt.to_period("M").dt.to_timestamp()
    monthly_rates = dict(zip(projection["month_start"], projection["inflation_monthly"]))

    values: list[dict[str, object]] = []
    current_value = float(start_value)
    for current_date in pd.date_range(start_date + pd.Timedelta(days=1), end_date, freq="D"):
        month_start = current_date.to_period("M").to_timestamp()
        monthly_rate = monthly_rates.get(month_start)
        if monthly_rate is None or pd.isna(monthly_rate):
            continue
        daily_factor = (1.0 + float(monthly_rate)) ** (1.0 / current_date.days_in_month)
        current_value *= daily_factor
        values.append(
            {
                "date": pd.Timestamp(current_date).normalize(),
                "band_upper": current_value,
                "source_type": "Proyección EcoGo",
            }
        )

    if not values:
        return pd.DataFrame(columns=["date", "band_upper", "source_type"])
    return pd.DataFrame(values)


def fetch_band_upper_series(reference_dates: pd.Series | pd.Index | list[object] | None = None) -> pd.DataFrame:
    try:
        frame = pd.read_excel(
            BCRA_BAND_XLSX_URL,
            sheet_name="Bandas de Flotacion Cambiaria",
            header=6,
        )
    except Exception:
        return pd.DataFrame(columns=["date", "band_upper", "source_type"])

    if "Fecha" not in frame.columns:
        return pd.DataFrame(columns=["date", "band_upper", "source_type"])
    upper_column = next((column for column in frame.columns if "Superior" in str(column)), None)
    if upper_column is None:
        return pd.DataFrame(columns=["date", "band_upper", "source_type"])

    series = frame.rename(columns={"Fecha": "date", upper_column: "band_upper"}).copy()
    series["date"] = pd.to_datetime(series["date"], errors="coerce")
    series["band_upper"] = pd.to_numeric(series["band_upper"], errors="coerce")
    series = series.dropna(subset=["date", "band_upper"])
    series["source_type"] = "BCRA oficial"

    derived_2025 = derive_2025_upper_band_series(reference_dates)
    combined = pd.concat([derived_2025, series[["date", "band_upper", "source_type"]]], ignore_index=True)
    combined = combined.dropna(subset=["date", "band_upper"])
    return combined.drop_duplicates(subset=["date"], keep="last").sort_values("date")


def build_upper_band_levels_series(
    history: pd.DataFrame,
    fx_futures_curve_latest: pd.DataFrame,
) -> list[dict[str, object]]:
    history_dates = pd.to_datetime(history["date"], errors="coerce") if "date" in history.columns else pd.Series(dtype="datetime64[ns]")
    futures_dates = pd.to_datetime(fx_futures_curve_latest["maturity_date"], errors="coerce") if "maturity_date" in fx_futures_curve_latest.columns else pd.Series(dtype="datetime64[ns]")

    valid_history_dates = history_dates.dropna()
    valid_futures_dates = futures_dates.dropna()
    if valid_history_dates.empty and valid_futures_dates.empty:
        return []

    end_candidates = [date for date in [valid_history_dates.max() if not valid_history_dates.empty else None, valid_futures_dates.max() if not valid_futures_dates.empty else None] if date is not None]
    end_date = max(end_candidates) if end_candidates else None
    if end_date is None:
        return []

    reference_dates = pd.date_range(SPREADS_START_DATE.normalize(), pd.Timestamp(end_date).normalize(), freq="D")
    base_series = fetch_band_upper_series(reference_dates)
    if base_series.empty:
        return []

    base_series = base_series.drop_duplicates(subset=["date"], keep="last").sort_values("date").copy()
    last_official_date = base_series.loc[base_series["source_type"] == "BCRA oficial", "date"].max()
    projection = pd.DataFrame(columns=["date", "band_upper", "source_type"])
    if not pd.isna(last_official_date) and pd.Timestamp(end_date).normalize() > pd.Timestamp(last_official_date).normalize():
        last_official_value = base_series.loc[base_series["date"] == last_official_date, "band_upper"].iloc[-1]
        projection = project_upper_band_with_ecogo(
            pd.Timestamp(last_official_date).normalize(),
            float(last_official_value),
            pd.Timestamp(end_date).normalize(),
            read_ecogo_inflation_projection(),
        )

    combined = pd.concat([base_series, projection], ignore_index=True)
    combined = combined.dropna(subset=["date", "band_upper"]).drop_duplicates(subset=["date"], keep="last").sort_values("date")
    return [
        {
            "date": row.date.strftime("%Y-%m-%d"),
            "upper_band": float_or_none(row.band_upper),
            "source_type": safe_str(row.source_type),
        }
        for row in combined.itertuples(index=False)
    ]


def build_fx_spreads(history: pd.DataFrame, fx_futures_curve_latest: pd.DataFrame | None = None) -> dict[str, object]:
    frame = history.copy()
    frame["date"] = pd.to_datetime(frame["date"], errors="coerce")
    frame["value_nominal"] = pd.to_numeric(frame["value_nominal"], errors="coerce")

    fx = frame[frame["asset"].isin(["usd_ars_ccl", "usd_ars_official"])].copy()
    if fx.empty:
        return {
            "ccl_vs_official": None,
            "official_vs_upper_band": None,
        }

    pivot = (
        fx.pivot_table(index="date", columns="asset", values="value_nominal", aggfunc="last")
        .reset_index()
        .sort_values("date")
    )
    pivot.loc[pivot["usd_ars_ccl"] <= 0, "usd_ars_ccl"] = pd.NA
    pivot.loc[pivot["usd_ars_official"] <= 0, "usd_ars_official"] = pd.NA

    ccl_official = None
    ccl_official_series: list[dict[str, object]] = []
    common = pivot.dropna(subset=["usd_ars_ccl", "usd_ars_official"]).copy()
    common = common[common["date"] >= SPREADS_START_DATE].copy()
    if not common.empty:
        row = common.iloc[-1]
        ccl_official = {
            "date": row["date"].strftime("%Y-%m-%d"),
            "ccl": float_or_none(row["usd_ars_ccl"]),
            "official": float_or_none(row["usd_ars_official"]),
            "spread_ars": float_or_none(row["usd_ars_ccl"] - row["usd_ars_official"]),
            "spread_pct": float_or_none((row["usd_ars_ccl"] / row["usd_ars_official"] - 1.0) * 100.0),
        }
        ccl_official_series = [
            {
                "date": entry.date.strftime("%Y-%m-%d"),
                "spread_pct": float_or_none((entry.usd_ars_ccl / entry.usd_ars_official - 1.0) * 100.0),
                "spread_ars": float_or_none(entry.usd_ars_ccl - entry.usd_ars_official),
                "official": float_or_none(entry.usd_ars_official),
                "ccl": float_or_none(entry.usd_ars_ccl),
            }
            for entry in common.itertuples(index=False)
        ]

    official_vs_upper_band = None
    official_vs_upper_band_series: list[dict[str, object]] = []
    official = pivot.dropna(subset=["usd_ars_official"])[["date", "usd_ars_official"]].copy()
    official = official[official["date"] >= SPREADS_START_DATE].copy()
    band_upper = fetch_band_upper_series(official["date"])
    if not official.empty and not band_upper.empty:
        merged = official.merge(band_upper, on="date", how="inner").sort_values("date")
        if not merged.empty:
            row = merged.iloc[-1]
            official_vs_upper_band = {
                "date": row["date"].strftime("%Y-%m-%d"),
                "official": float_or_none(row["usd_ars_official"]),
                "upper_band": float_or_none(row["band_upper"]),
                "spread_ars": float_or_none(row["band_upper"] - row["usd_ars_official"]),
                "spread_pct": float_or_none((row["band_upper"] / row["usd_ars_official"] - 1.0) * 100.0),
                "source_url": BCRA_BAND_XLSX_URL,
                "source_note": "BCRA XLSX oficial desde 2026-01-02; 2025 reconstruido con la regla oficial de +1% mensual del régimen de bandas.",
            }
            official_vs_upper_band_series = [
                {
                    "date": entry.date.strftime("%Y-%m-%d"),
                    "spread_pct": float_or_none((entry.band_upper / entry.usd_ars_official - 1.0) * 100.0),
                    "spread_ars": float_or_none(entry.band_upper - entry.usd_ars_official),
                    "official": float_or_none(entry.usd_ars_official),
                    "upper_band": float_or_none(entry.band_upper),
                }
                for entry in merged.itertuples(index=False)
            ]

    return {
        "ccl_vs_official": ccl_official,
        "ccl_vs_official_series": ccl_official_series,
        "official_vs_upper_band": official_vs_upper_band,
        "official_vs_upper_band_series": official_vs_upper_band_series,
        "upper_band_levels_series": build_upper_band_levels_series(
            frame,
            fx_futures_curve_latest if fx_futures_curve_latest is not None else pd.DataFrame(),
        ),
    }


def build_macro_compare(lecap_macro_history: pd.DataFrame, lecap_macro_monthly: pd.DataFrame) -> list[dict[str, object]]:
    if lecap_macro_monthly.empty:
        return []

    monthly = lecap_macro_monthly.copy()
    monthly["month"] = pd.to_datetime(monthly["month"], errors="coerce")
    monthly["inflation_monthly_pct"] = pd.to_numeric(monthly["inflation_monthly_pct"], errors="coerce")
    monthly["a3500_devaluation_pct"] = pd.to_numeric(monthly["a3500_devaluation_pct"], errors="coerce")

    avg_tasa = pd.DataFrame(columns=["month", "average_tem_pct"])
    if not lecap_macro_history.empty:
        history = lecap_macro_history.copy()
        history["date"] = pd.to_datetime(history["date"], errors="coerce")
        history["average_tem_pct"] = pd.to_numeric(history["average_tem_pct"], errors="coerce")
        history = history.dropna(subset=["date", "average_tem_pct"]).copy()
        history["month"] = history["date"].dt.to_period("M").dt.to_timestamp()
        avg_tasa = (
            history.groupby("month", as_index=False)["average_tem_pct"]
            .mean()
            .sort_values("month")
        )

    merged = monthly.merge(avg_tasa, on="month", how="left").sort_values("month")
    return [
        {
            "month": row.month.strftime("%Y-%m-%d") if not pd.isna(row.month) else None,
            "month_label": safe_str(row.month_label),
            "average_tem_pct": float_or_none(row.average_tem_pct),
            "inflation_monthly_pct": float_or_none(row.inflation_monthly_pct),
            "devaluation_monthly_pct": float_or_none(row.a3500_devaluation_pct),
        }
        for row in merged.itertuples(index=False)
    ]


def build_macro_compare_daily(lecap_macro_history: pd.DataFrame, lecap_macro_monthly: pd.DataFrame) -> dict[str, list[dict[str, object]]]:
    daily_series: list[dict[str, object]] = []
    if not lecap_macro_history.empty:
        daily_average = (
            lecap_macro_history.copy()[["date", "average_tem_pct", "average_price", "symbol_count"]]
            .drop_duplicates(subset=["date"])
            .sort_values("date")
        )
        daily_average["date"] = pd.to_datetime(daily_average["date"], errors="coerce")
        daily_average["average_tem_pct"] = pd.to_numeric(daily_average["average_tem_pct"], errors="coerce")
        daily_average["average_price"] = pd.to_numeric(daily_average["average_price"], errors="coerce")
        daily_average["symbol_count"] = pd.to_numeric(daily_average["symbol_count"], errors="coerce")
        daily_average = daily_average.dropna(subset=["date", "average_tem_pct"]).copy()
        daily_series = [
            {
                "date": row.date.strftime("%Y-%m-%d"),
                "average_tem_pct": float_or_none(row.average_tem_pct),
                "average_price": float_or_none(row.average_price),
                "symbol_count": int(row.symbol_count) if not pd.isna(row.symbol_count) else None,
            }
            for row in daily_average.itertuples(index=False)
        ]

    monthly_inflation: list[dict[str, object]] = []
    monthly_devaluation: list[dict[str, object]] = []
    observed_inflation_last_6m: list[dict[str, object]] = []
    if not lecap_macro_monthly.empty:
        monthly = lecap_macro_monthly.copy()
        monthly["month"] = pd.to_datetime(monthly["month"], errors="coerce")
        monthly["plot_date"] = monthly["month"] + pd.offsets.MonthEnd(0)
        monthly["inflation_monthly_pct"] = pd.to_numeric(monthly["inflation_monthly_pct"], errors="coerce")
        monthly["a3500_devaluation_pct"] = pd.to_numeric(monthly["a3500_devaluation_pct"], errors="coerce")
        monthly = monthly.sort_values("month")
        monthly_inflation = [
            {
                "month": row.month.strftime("%Y-%m-%d"),
                "plot_date": row.plot_date.strftime("%Y-%m-%d"),
                "month_label": safe_str(row.month_label),
                "value_pct": float_or_none(row.inflation_monthly_pct),
            }
            for row in monthly.dropna(subset=["month", "inflation_monthly_pct"]).itertuples(index=False)
        ]
        monthly_devaluation = [
            {
                "month": row.month.strftime("%Y-%m-%d"),
                "plot_date": row.plot_date.strftime("%Y-%m-%d"),
                "month_label": safe_str(row.month_label),
                "value_pct": float_or_none(row.a3500_devaluation_pct),
            }
            for row in monthly.dropna(subset=["month", "a3500_devaluation_pct"]).itertuples(index=False)
        ]
        observed_rows = monthly.dropna(subset=["month", "inflation_monthly_pct"]).tail(6)
        observed_inflation_last_6m = [
            {
                "month": row.month.strftime("%Y-%m-%d"),
                "plot_date": row.plot_date.strftime("%Y-%m-%d"),
                "month_label": safe_str(row.month_label),
                "value_pct": float_or_none(row.inflation_monthly_pct),
            }
            for row in observed_rows.itertuples(index=False)
        ]

    return {
        "daily_average": daily_series,
        "monthly_inflation": monthly_inflation,
        "monthly_devaluation": monthly_devaluation,
        "observed_inflation_last_6m": observed_inflation_last_6m,
    }


def build_payload() -> dict[str, object]:
    latest = read_csv(LATEST_FILE)
    history = read_csv(SERIES_FILE)
    lecap_latest = read_csv(LECAP_FILE)
    lecap_history = read_csv(LECAP_HISTORY_FILE)
    cer_curve_latest = read_csv(CER_CURVE_FILE)
    cer_curve_history = read_csv(CER_CURVE_HISTORY_FILE)
    hard_dollar_latest = read_csv(HARD_DOLLAR_CURVE_FILE)
    hard_dollar_history = read_csv(HARD_DOLLAR_CURVE_HISTORY_FILE)
    hard_dollar_weighted_parity_history = read_csv(HARD_DOLLAR_WEIGHTED_PARITY_HISTORY_FILE)
    fx_futures_curve_latest = read_csv(FX_FUTURES_CURVE_LATEST_FILE)
    dollar_linked_history = read_csv(DOLLAR_LINKED_HISTORY_FILE)
    dollar_linked_latest = read_csv(DOLLAR_LINKED_LATEST_FILE)
    monetary_variables_latest = read_csv(MONETARY_VARIABLES_LATEST_FILE)
    monetary_variables_economy_pib_history = read_csv(MONETARY_VARIABLES_ECONOMY_PIB_HISTORY_FILE)
    monetary_variables_peso_usd_pib_history = read_csv(MONETARY_VARIABLES_PESO_USD_PIB_HISTORY_FILE)
    monetary_variables_usd_stocks_history = read_csv(MONETARY_VARIABLES_USD_STOCKS_HISTORY_FILE)
    monetary_variables_block_freshness = read_csv(MONETARY_VARIABLES_BLOCK_FRESHNESS_FILE)
    monetary_variables_source_table = read_csv(MONETARY_VARIABLES_SOURCE_TABLE_FILE)
    monetary_rates_tna_series = read_csv(MONETARY_RATES_TNA_SERIES_FILE)
    monetary_rates_tna_reference = read_csv(MONETARY_RATES_TNA_REFERENCE_FILE)
    monetary_rates_tna_latest = read_csv(MONETARY_RATES_TNA_LATEST_FILE)
    fci_category_history = read_csv(FCI_CATEGORY_HISTORY_FILE)
    fci_currency_history = read_csv(FCI_CURRENCY_HISTORY_FILE)
    fci_horizon_history = read_csv(FCI_HORIZON_HISTORY_FILE)
    fci_money_market_latest = read_csv(FCI_MONEY_MARKET_LATEST_FILE)
    fci_reported_slice_history = read_csv(FCI_REPORTED_SLICE_HISTORY_FILE)
    breakeven = read_csv(BREAKEVEN_FILE)
    lecap_macro_history = read_csv(LECAP_MACRO_HISTORY_FILE)
    lecap_macro_monthly = read_csv(LECAP_MACRO_MONTHLY_FILE)
    alerts = read_csv(ALERTS_FILE)
    meta = json.loads(META_FILE.read_text(encoding="utf-8")) if META_FILE.exists() else {}
    freshness = build_freshness(
        latest,
        lecap_history,
        cer_curve_history,
        hard_dollar_history,
        fx_futures_curve_latest,
        dollar_linked_history,
        monetary_variables_latest,
        fci_category_history,
        fci_currency_history,
        fci_horizon_history,
        fci_money_market_latest,
        fci_reported_slice_history,
    )

    return {
        "meta": meta,
        "freshness": freshness,
        "overview": build_overview(latest),
        "hero_metrics": build_hero_metrics(latest),
        "hero_series": build_series(history),
        "external_context": build_external_context(latest, history),
        "market_asset_series": build_market_asset_series(history),
        "fixed_curve": build_fixed_curve(lecap_latest),
        "fixed_curve_history": build_fixed_curve_history(lecap_history),
        "cer_curve": build_cer_curve(cer_curve_latest),
        "cer_curve_history": build_cer_curve_history(cer_curve_history),
        "hard_dollar_curve": build_hard_dollar_curve(hard_dollar_latest),
        "hard_dollar_curve_history": build_hard_dollar_curve_history(hard_dollar_history),
        "hard_dollar_weighted_parity": build_hard_dollar_weighted_parity(hard_dollar_weighted_parity_history),
        "hard_dollar_country_risk": build_hard_dollar_country_risk(history),
        "fx_futures_curve": build_fx_futures_curve(fx_futures_curve_latest),
        "dollar_linked": build_dollar_linked_monitor(dollar_linked_history, dollar_linked_latest),
        "monetary_variables": build_monetary_variables(
            monetary_variables_latest,
            monetary_variables_economy_pib_history,
            monetary_variables_peso_usd_pib_history,
            monetary_variables_usd_stocks_history,
            monetary_variables_block_freshness,
            monetary_variables_source_table,
            monetary_rates_tna_series,
            monetary_rates_tna_reference,
            monetary_rates_tna_latest,
        ),
        "fci_category_patrimony": build_fci_category_history(fci_category_history),
        "fci_currency_patrimony": build_fci_currency_history(fci_currency_history),
        "fci_horizon_patrimony": build_fci_horizon_history(fci_horizon_history),
        "fci_aum_matrix": build_fci_aum_matrix(fci_category_history, fci_currency_history, fci_horizon_history),
        "fci_reported_slices": build_fci_reported_slices(fci_reported_slice_history, fci_category_history),
        "fci_money_market": build_fci_money_market_latest(fci_money_market_latest),
        "breakeven_curve": build_breakeven_curve(breakeven, lecap_macro_monthly),
        "macro_compare": build_macro_compare(lecap_macro_history, lecap_macro_monthly),
        "macro_compare_daily": build_macro_compare_daily(lecap_macro_history, lecap_macro_monthly),
        "fx_spreads": build_fx_spreads(history, fx_futures_curve_latest),
        "alerts": build_alerts(alerts),
        "hard_dollar": build_hard_dollar(latest, hard_dollar_latest),
    }


def main() -> None:
    WEB_DATA_DIR.mkdir(parents=True, exist_ok=True)
    payload = build_payload()
    serialized = json.dumps(payload, indent=2, ensure_ascii=False, allow_nan=False)
    OUTPUT_FILE.write_text(serialized, encoding="utf-8")
    OUTPUT_JS_FILE.write_text(f"window.__DASHBOARD_PAYLOAD__ = {serialized};\n", encoding="utf-8")
    print(f"Wrote {OUTPUT_FILE}")
    print(f"Wrote {OUTPUT_JS_FILE}")


if __name__ == "__main__":
    main()
