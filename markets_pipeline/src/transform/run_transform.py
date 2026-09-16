from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from src.common.cer_coupon_curve import build_cer_coupon_metrics
from src.common.cer_curve import CER_ZERO_COUPON_COLUMNS, build_cer_zero_coupon_metrics
from src.common.a3_futures import A3_FUTURES_COLUMNS, fetch_a3_dollar_futures_curve
from src.common.argentinadatos_fci import (
    FCI_CATEGORY_COLUMNS,
    FCI_CURRENCY_COLUMNS,
    FCI_HORIZON_COLUMNS,
    FCI_MONEY_MARKET_COLUMNS,
    FCI_REPORTED_SLICE_COLUMNS,
    fetch_fci_category_history,
    fetch_fci_currency_history,
    fetch_fci_horizon_history,
    fetch_fci_reported_slice_history,
    fetch_money_market_fund_latest,
)
from src.common.hard_dollar_curve import HARD_DOLLAR_COLUMNS, build_hard_dollar_metrics
from src.common.hard_dollar_weighted_parity import (
    HARD_DOLLAR_WEIGHTED_PARITY_COLUMNS,
    HARD_DOLLAR_WEIGHTED_PARITY_LATEST_COLUMNS,
    build_weighted_parity_history_from_frames,
)
from src.common.dollar_linked import (
    DOLLAR_LINKED_HISTORY_COLUMNS,
    DOLLAR_LINKED_HISTORY_WINDOW,
    DOLLAR_LINKED_LATEST_COLUMNS,
    DOLLAR_LINKED_SNAPSHOT_COLUMNS,
    DOLLAR_LINKED_SYMBOLS,
)
from src.common.fixed_income import DEFAULT_IOL_INSTRUMENTS_DROP_PATH, DEFAULT_IOL_INSTRUMENTS_METADATA_PATH
from src.common.io import PROCESSED_DIR, RAW_DIR, ensure_data_dirs
from src.common.monetary_variables import (
    MONETARY_BLOCK_FRESHNESS_COLUMNS,
    MONETARY_BLOCK_WINDOWS,
    MONETARY_CORE_METRICS,
    MONETARY_ECONOMY_PIB_HISTORY_COLUMNS,
    MONETARY_PESO_USD_PIB_HISTORY_COLUMNS,
    MONETARY_RATE_UNIT,
    MONETARY_RATES_TNA_LABELS,
    MONETARY_RATES_TNA_LATEST_COLUMNS,
    MONETARY_RATES_TNA_METRICS,
    MONETARY_RATES_TNA_RAW_COLUMNS,
    MONETARY_RATES_TNA_REFERENCE_COLUMNS,
    MONETARY_RATES_TNA_SERIES_COLUMNS,
    MONETARY_RATES_TNA_SPECS,
    MONETARY_SOURCE_TABLE_COLUMNS,
    MONETARY_USD_STOCKS_HISTORY_COLUMNS,
    MONETARY_VARIABLE_CORE_LATEST_COLUMNS,
    MONETARY_VARIABLE_LABELS,
    MONETARY_VARIABLE_METRICS,
    MONETARY_VARIABLE_UNIT,
    MONETARY_VARIABLE_SPECS,
    MONETARY_VARIABLE_WINDOW_MONTHS,
)
from src.ingest.iol import load_iol_instrument_metadata
from src.common.cer_breakeven import (
    build_breakeven_forward_monthly,
    build_breakeven_monitor_live,
    load_cer_index_history,
)


RAW_PRICES_FILE = RAW_DIR / "prices_nominal_daily.csv"
RAW_CPI_FILE = RAW_DIR / "cpi_monthly.csv"
RAW_GDP_FILE = RAW_DIR / "indec_gdp_nominal_quarterly.csv"
ROOT = Path(__file__).resolve().parents[2]
ECOGO_INFLATION_PROJECTION_FILE = ROOT / "legacy" / "ecogo_projection_inputs" / "proy_infla_ecogo.xlsx"
RAW_FUNDAMENTALS_FILE = RAW_DIR / "fundamentals_annual_long.csv"
RAW_WARNINGS_FILE = RAW_DIR / "fundamentals_warnings.txt"
IOL_ZERO_COUPON_METRICS_FILE = PROCESSED_DIR / "iol_zero_coupon_metrics.csv"
CER_BOND_HISTORY_FILE = RAW_DIR / "cer_bond_prices_history.csv"
CER_INDEX_FILE = RAW_DIR / "bcra_cer_index_daily.csv"
IOL_MARKET_PRICES_FILE = RAW_DIR / "iol_market_prices_history.csv"
DOLLAR_LINKED_SNAPSHOT_FILE = RAW_DIR / "iol_dollar_linked_snapshot.csv"
MONETARY_VARIABLES_RAW_FILE = RAW_DIR / "bcra_monetary_variables_daily.csv"
MONETARY_RATES_TNA_RAW_FILE = RAW_DIR / "bcra_rates_tna_daily.csv"
HARD_DOLLAR_REFERENCE_METADATA_FILE = RAW_DIR / "hard_dollar_reference_metadata.csv"
HARD_DOLLAR_REFERENCE_FLOW_FILE = RAW_DIR / "hard_dollar_reference_flows.csv"

SERIES_FILE = PROCESSED_DIR / "market_series_daily.csv"
LATEST_FILE = PROCESSED_DIR / "latest_snapshot.csv"
FUNDAMENTALS_FILE = PROCESSED_DIR / "fundamentals_annual.csv"
FUNDAMENTALS_LATEST_FILE = PROCESSED_DIR / "fundamentals_latest.csv"
ALERTS_FILE = PROCESSED_DIR / "alerts_latest.csv"
META_FILE = PROCESSED_DIR / "pipeline_metadata.json"
LECAP_MONITOR_HISTORY_FILE = PROCESSED_DIR / "lecap_monitor_history.csv"
LECAP_MONITOR_LATEST_FILE = PROCESSED_DIR / "lecap_monitor_latest.csv"
BREAKEVEN_HISTORY_FILE = PROCESSED_DIR / "cer_breakeven_history.csv"
BREAKEVEN_LATEST_FILE = PROCESSED_DIR / "cer_breakeven_latest.csv"
BREAKEVEN_MONTHLY_HISTORY_FILE = PROCESSED_DIR / "cer_breakeven_monthly_history.csv"
BREAKEVEN_MONTHLY_LATEST_FILE = PROCESSED_DIR / "cer_breakeven_monthly_latest.csv"
OBSERVED_INFLATION_FILE = PROCESSED_DIR / "inflation_observed_monthly.csv"
LECAP_MACRO_HISTORY_FILE = PROCESSED_DIR / "lecap_macro_history.csv"
LECAP_MACRO_MONTHLY_FILE = PROCESSED_DIR / "lecap_macro_monthly.csv"
CER_CURVE_HISTORY_FILE = PROCESSED_DIR / "cer_curve_history.csv"
CER_CURVE_LATEST_FILE = PROCESSED_DIR / "cer_curve_latest.csv"
HARD_DOLLAR_CURVE_HISTORY_FILE = PROCESSED_DIR / "hard_dollar_curve_history.csv"
HARD_DOLLAR_CURVE_LATEST_FILE = PROCESSED_DIR / "hard_dollar_curve_latest.csv"
HARD_DOLLAR_AUDIT_FILE = PROCESSED_DIR / "hard_dollar_cashflow_audit.csv"
HARD_DOLLAR_WEIGHTED_PARITY_HISTORY_FILE = PROCESSED_DIR / "hard_dollar_weighted_parity_history.csv"
HARD_DOLLAR_WEIGHTED_PARITY_LATEST_FILE = PROCESSED_DIR / "hard_dollar_weighted_parity_latest.csv"
FX_FUTURES_CURVE_LATEST_FILE = PROCESSED_DIR / "fx_futures_curve_latest.csv"
DOLLAR_LINKED_HISTORY_FILE = PROCESSED_DIR / "dollar_linked_volume_history.csv"
DOLLAR_LINKED_LATEST_FILE = PROCESSED_DIR / "dollar_linked_volume_latest.csv"
MONETARY_VARIABLES_LATEST_FILE = PROCESSED_DIR / "monetary_variables_latest.csv"
MONETARY_VARIABLES_ECONOMY_PIB_HISTORY_FILE = PROCESSED_DIR / "monetary_variables_economy_pib_history.csv"
MONETARY_VARIABLES_PESO_USD_PIB_HISTORY_FILE = PROCESSED_DIR / "monetary_variables_peso_usd_pib_history.csv"
MONETARY_VARIABLES_USD_STOCKS_HISTORY_FILE = PROCESSED_DIR / "monetary_variables_usd_stocks_history.csv"
MONETARY_VARIABLES_BLOCK_FRESHNESS_FILE = PROCESSED_DIR / "monetary_variables_block_freshness.csv"
MONETARY_VARIABLES_SOURCE_TABLE_FILE = PROCESSED_DIR / "monetary_variables_source_table.csv"
MONETARY_RATES_TNA_SERIES_FILE = PROCESSED_DIR / "monetary_rates_tna_series.csv"
MONETARY_RATES_TNA_REFERENCE_FILE = PROCESSED_DIR / "monetary_rates_tna_reference.csv"
MONETARY_RATES_TNA_LATEST_FILE = PROCESSED_DIR / "monetary_rates_tna_latest.csv"
FCI_CATEGORY_HISTORY_FILE = PROCESSED_DIR / "fci_category_patrimony_history.csv"
FCI_CURRENCY_HISTORY_FILE = PROCESSED_DIR / "fci_currency_patrimony_history.csv"
FCI_HORIZON_HISTORY_FILE = PROCESSED_DIR / "fci_horizon_patrimony_history.csv"
FCI_MONEY_MARKET_LATEST_FILE = PROCESSED_DIR / "fci_money_market_latest.csv"
FCI_REPORTED_SLICE_HISTORY_FILE = PROCESSED_DIR / "fci_reported_slice_history.csv"
FCI_FETCH_LOOKBACK_DAYS = 45
FCI_FETCH_MAX_WORKERS = 8
FCI_FETCH_TIMEOUT_SECONDS = 8
ALERT_COLUMNS = [
    "as_of_date",
    "category",
    "asset",
    "asset_label",
    "severity",
    "metric_key",
    "metric_label",
    "observed_value",
    "threshold_value",
    "direction",
]
LECAP_MONITOR_COLUMNS = [
    "date",
    "symbol",
    "instrument_family",
    "maturity_date",
    "next_business_day",
    "payment_date",
    "days_360_to_maturity",
    "tea_pct",
    "tem_pct",
    "modified_duration",
    "price",
    "technical_price",
    "parity",
    "volume",
    "volume_billions",
]
LECAP_MACRO_HISTORY_COLUMNS = [
    "date",
    "symbol",
    "price",
    "tem_pct",
    "tea_pct",
    "average_tem_pct",
    "average_price",
    "symbol_count",
]
LECAP_MACRO_MONTHLY_COLUMNS = [
    "month",
    "month_label",
    "inflation_monthly_pct",
    "a3500_avg",
    "a3500_devaluation_pct",
]


def _apply_performance_basis(frame: pd.DataFrame) -> pd.DataFrame:
    frame = frame.sort_values(["date", "category", "asset_label"]).reset_index(drop=True).copy()
    frame["performance_basis"] = "real"
    frame["value_metric_today"] = frame["value_real_today"]
    external_mask = frame["category"].eq("External")
    frame.loc[external_mask, "performance_basis"] = "nominal"
    frame.loc[external_mask, "value_metric_today"] = pd.to_numeric(
        frame.loc[external_mask, "value_nominal"],
        errors="coerce",
    )

    ccl_reference = (
        frame[frame["asset"] == "usd_ars_ccl"][["date", "value_nominal"]]
        .dropna(subset=["date", "value_nominal"])
        .sort_values("date")
        .drop_duplicates(subset=["date"], keep="last")
        .rename(columns={"value_nominal": "usd_ars_ccl"})
    )
    if ccl_reference.empty:
        return frame

    enriched = pd.merge_asof(
        frame.sort_values("date"),
        ccl_reference,
        on="date",
        direction="backward",
        allow_exact_matches=True,
    )
    enriched["usd_ars_ccl"] = pd.to_numeric(enriched["usd_ars_ccl"], errors="coerce")

    equities_mask = enriched["category"].eq("Equities")
    merval_mask = equities_mask & enriched["asset"].eq("merval_usd")
    ccl_equity_mask = equities_mask & ~merval_mask
    valid_ccl_mask = ccl_equity_mask & enriched["usd_ars_ccl"].notna() & (enriched["usd_ars_ccl"] > 0)
    external_mask = enriched["category"].eq("External")

    enriched.loc[merval_mask, "performance_basis"] = "ccl"
    enriched.loc[merval_mask, "value_metric_today"] = pd.to_numeric(
        enriched.loc[merval_mask, "value_nominal"],
        errors="coerce",
    )
    enriched.loc[valid_ccl_mask, "performance_basis"] = "ccl"
    enriched.loc[valid_ccl_mask, "value_metric_today"] = (
        pd.to_numeric(enriched.loc[valid_ccl_mask, "value_nominal"], errors="coerce")
        / enriched.loc[valid_ccl_mask, "usd_ars_ccl"]
    )
    enriched.loc[external_mask, "performance_basis"] = "nominal"
    enriched.loc[external_mask, "value_metric_today"] = pd.to_numeric(
        enriched.loc[external_mask, "value_nominal"],
        errors="coerce",
    )

    return enriched.drop(columns=["usd_ars_ccl"])


def _compute_group_metrics(group: pd.DataFrame) -> pd.DataFrame:
    group = group.sort_values("date").copy()
    metric_columns = ["ret_1d", "ret_1m", "real_mean_252", "real_std_252", "zscore_252", "index_base_100", "ret_ytd"]
    for column in metric_columns:
        group[column] = pd.NA

    valid_mask = group["value_metric_today"].notna()
    if not valid_mask.any():
        return group

    valid = group.loc[valid_mask, ["date", "value_metric_today"]].copy()
    valid["ret_1d"] = valid["value_metric_today"].pct_change(1, fill_method=None)
    valid["ret_1m"] = valid["value_metric_today"].pct_change(21, fill_method=None)
    valid["real_mean_252"] = valid["value_metric_today"].rolling(252, min_periods=60).mean()
    valid["real_std_252"] = valid["value_metric_today"].rolling(252, min_periods=60).std()
    valid["zscore_252"] = (valid["value_metric_today"] - valid["real_mean_252"]) / valid["real_std_252"]

    first_value = valid["value_metric_today"].iloc[0]
    valid["index_base_100"] = valid["value_metric_today"] / first_value * 100.0

    valid["ret_ytd"] = pd.NA
    for _, year_group in valid.groupby(valid["date"].dt.year):
        base = year_group["value_metric_today"].iloc[0]
        valid.loc[year_group.index, "ret_ytd"] = year_group["value_metric_today"] / base - 1.0

    group.loc[valid.index, metric_columns] = valid[metric_columns]
    return group


def _build_alerts(latest: pd.DataFrame) -> pd.DataFrame:
    rows: list[dict[str, object]] = []
    filtered = latest[latest["category"] != "External"].copy()
    for row in filtered.itertuples(index=False):
        if pd.notna(row.zscore_252) and row.zscore_252 >= 1.5:
            rows.append(
                {
                    "as_of_date": row.date.date().isoformat(),
                    "category": row.category,
                    "asset": row.asset,
                    "asset_label": row.asset_label,
                    "severity": "warning",
                    "metric_key": "zscore_252",
                    "metric_label": "Z-score 252d",
                    "observed_value": round(float(row.zscore_252), 3),
                    "threshold_value": 1.5,
                    "direction": "above",
                }
            )
        if pd.notna(row.zscore_252) and row.zscore_252 <= -1.5:
            rows.append(
                {
                    "as_of_date": row.date.date().isoformat(),
                    "category": row.category,
                    "asset": row.asset,
                    "asset_label": row.asset_label,
                    "severity": "warning",
                    "metric_key": "zscore_252",
                    "metric_label": "Z-score 252d",
                    "observed_value": round(float(row.zscore_252), 3),
                    "threshold_value": -1.5,
                    "direction": "below",
                }
            )
        if pd.notna(row.ret_1m) and abs(float(row.ret_1m)) >= 0.10:
            rows.append(
                {
                    "as_of_date": row.date.date().isoformat(),
                    "category": row.category,
                    "asset": row.asset,
                    "asset_label": row.asset_label,
                    "severity": "info" if abs(float(row.ret_1m)) < 0.18 else "critical",
                    "metric_key": "ret_1m_pct",
                    "metric_label": "Retorno 1M",
                    "observed_value": round(float(row.ret_1m) * 100.0, 3),
                    "threshold_value": 10.0,
                    "direction": "absolute",
                }
            )
    return pd.DataFrame(rows, columns=ALERT_COLUMNS)


def _metadata_note() -> str:
    if not RAW_WARNINGS_FILE.exists():
        return "Processed dataset generated from configured sources."

    warning_text = RAW_WARNINGS_FILE.read_text(encoding="utf-8")
    has_bcra = "kind=bcra_" in warning_text
    has_byma = "kind=byma_" in warning_text
    has_iol = "kind=iol_" in warning_text
    has_synthetic = "kind=synthetic_" in warning_text
    if has_bcra and has_byma and has_iol and has_synthetic:
        return "Mixed dataset: BCRA, IOL, and BYMA live sources for selected series, scaffold placeholders for the remaining assets."
    if has_bcra and has_byma and has_iol:
        return "Mixed dataset: BCRA, IOL, and BYMA live sources for selected series."
    if has_bcra and has_iol and has_synthetic:
        return "Mixed dataset: BCRA and IOL live sources for selected series, scaffold placeholders for the remaining assets."
    if has_byma and has_iol and has_synthetic:
        return "Mixed dataset: IOL and BYMA live sources for selected listed assets, scaffold placeholders for the remaining assets."
    if has_bcra and has_iol:
        return "Mixed dataset: BCRA and IOL live sources for selected series."
    if has_byma and has_iol:
        return "Mixed dataset: IOL and BYMA live sources for selected listed assets."
    if has_bcra and has_byma and has_synthetic:
        return "Mixed dataset: BCRA and BYMA live sources for selected series, scaffold placeholders for the remaining assets."
    if has_bcra and has_byma:
        return "Mixed dataset: BCRA and BYMA live sources for selected series."
    if has_byma and has_synthetic:
        return "Mixed dataset: BYMA live sources for selected listed assets, scaffold placeholders for the remaining assets."
    if has_iol and has_synthetic:
        return "Mixed dataset: IOL live sources for selected listed assets, scaffold placeholders for the remaining assets."
    if has_bcra:
        return "Dataset generated from live BCRA sources."
    if has_byma:
        return "Dataset generated from live BYMA sources."
    if has_iol:
        return "Dataset generated from live IOL sources."
    if has_synthetic:
        return "Synthetic scaffold dataset. Replace raw ingestion with production connectors."
    return "Processed dataset generated from configured sources."


def _build_lecap_monitor() -> tuple[pd.DataFrame, pd.DataFrame]:
    if not IOL_ZERO_COUPON_METRICS_FILE.exists() or IOL_ZERO_COUPON_METRICS_FILE.stat().st_size == 0:
        empty = pd.DataFrame(columns=LECAP_MONITOR_COLUMNS)
        return empty.copy(), empty.copy()

    metrics = pd.read_csv(
        IOL_ZERO_COUPON_METRICS_FILE,
        parse_dates=["date", "issue_date", "maturity_date", "next_business_day", "payment_date"],
    )
    fixed_rate = metrics[metrics["instrument_family"].astype(str).str.lower().isin(["lecap", "boncap"])].copy()
    if fixed_rate.empty:
        empty = pd.DataFrame(columns=LECAP_MONITOR_COLUMNS)
        return empty.copy(), empty.copy()

    fixed_rate = fixed_rate.dropna(subset=["date", "maturity_date", "tea", "modified_duration"]).copy()
    fixed_rate = fixed_rate[fixed_rate["days_360_to_maturity"] > 0].copy()
    if fixed_rate.empty:
        empty = pd.DataFrame(columns=LECAP_MONITOR_COLUMNS)
        return empty.copy(), empty.copy()

    latest = (
        fixed_rate.sort_values("date")
        .groupby("symbol", as_index=False)
        .tail(1)
        .reset_index(drop=True)
    )
    reference_date = fixed_rate["date"].max()
    active_cutoff = max(pd.Timestamp(reference_date).normalize(), pd.Timestamp.today().normalize())
    latest = latest[pd.to_datetime(latest["payment_date"], errors="coerce") > active_cutoff].copy()
    if latest.empty:
        empty = pd.DataFrame(columns=LECAP_MONITOR_COLUMNS)
        return empty.copy(), empty.copy()

    active_symbols = latest["symbol"].dropna().astype(str).str.upper().unique().tolist()
    history = fixed_rate[fixed_rate["symbol"].astype(str).str.upper().isin(active_symbols)].copy()

    for frame in (history, latest):
        frame["symbol"] = frame["symbol"].astype(str).str.upper()
        frame["instrument_family"] = frame["instrument_family"].astype(str).str.lower()
        frame["tea_pct"] = pd.to_numeric(frame["tea"], errors="coerce") * 100.0
        frame["tem_pct"] = pd.to_numeric(frame["tem"], errors="coerce") * 100.0
        frame["modified_duration"] = pd.to_numeric(frame["modified_duration"], errors="coerce")
        frame["days_360_to_maturity"] = pd.to_numeric(frame["days_360_to_maturity"], errors="coerce")
        frame["price"] = pd.to_numeric(frame["price"], errors="coerce")
        frame["technical_price"] = pd.to_numeric(frame["technical_price"], errors="coerce")
        frame["parity"] = pd.to_numeric(frame["parity"], errors="coerce")
        frame["volume"] = pd.to_numeric(frame["volume"], errors="coerce")
        frame["volume_billions"] = frame["volume"] / 1_000_000_000.0

    history = (
        history[LECAP_MONITOR_COLUMNS]
        .sort_values(["date", "days_360_to_maturity", "instrument_family", "symbol"])
        .reset_index(drop=True)
    )
    latest = (
        latest[LECAP_MONITOR_COLUMNS]
        .sort_values(["days_360_to_maturity", "instrument_family", "symbol"])
        .reset_index(drop=True)
    )
    return history, latest


def _build_observed_inflation(cpi: pd.DataFrame) -> pd.DataFrame:
    observed = cpi.sort_values("month").copy()
    observed["inflation_monthly"] = observed["cpi_index"].pct_change()
    observed["inflation_monthly_pct"] = observed["inflation_monthly"] * 100.0
    observed["month_label"] = observed["month"].dt.strftime("%Y-%m")
    observed = observed.dropna(subset=["inflation_monthly"]).reset_index(drop=True)
    return observed[["month", "month_label", "cpi_index", "inflation_monthly", "inflation_monthly_pct"]]


def _read_ecogo_inflation_projection() -> pd.DataFrame:
    if not ECOGO_INFLATION_PROJECTION_FILE.exists():
        return pd.DataFrame(columns=["month", "inflation_monthly"])

    workbook = load_workbook(ECOGO_INFLATION_PROJECTION_FILE, data_only=True, read_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    workbook.close()

    frame = pd.DataFrame(rows, columns=["month", "inflation_monthly"])
    frame["month"] = pd.to_datetime(frame["month"], errors="coerce").dt.to_period("M").dt.to_timestamp()
    frame["inflation_monthly"] = pd.to_numeric(frame["inflation_monthly"], errors="coerce")
    frame = frame.dropna(subset=["month", "inflation_monthly"]).copy()
    return frame.sort_values("month").drop_duplicates(subset=["month"], keep="last").reset_index(drop=True)


def _annualize_monthly_inflation(monthly_rate: pd.Series) -> pd.Series:
    monthly = pd.to_numeric(monthly_rate, errors="coerce")
    return ((1.0 + monthly) ** 12 - 1.0) * 100.0


def _expand_monthly_reference_daily(
    frame: pd.DataFrame,
    *,
    metric: str,
    label: str,
    source_type: str,
) -> pd.DataFrame:
    if frame.empty:
        return pd.DataFrame(columns=MONETARY_RATES_TNA_REFERENCE_COLUMNS)

    rows: list[dict[str, object]] = []
    local = frame.copy()
    local["month"] = pd.to_datetime(local["month"], errors="coerce").dt.to_period("M").dt.to_timestamp()
    local["value"] = pd.to_numeric(local["value"], errors="coerce")
    local = local.dropna(subset=["month", "value"]).sort_values("month").reset_index(drop=True)
    if local.empty:
        return pd.DataFrame(columns=MONETARY_RATES_TNA_REFERENCE_COLUMNS)

    for row in local.itertuples(index=False):
        month_start = pd.Timestamp(row.month).normalize()
        month_end = month_start + pd.offsets.MonthEnd(0)
        for date in pd.date_range(month_start, month_end, freq="D"):
            rows.append(
                {
                    "date": pd.Timestamp(date).normalize(),
                    "metric": metric,
                    "label": label,
                    "value": float(row.value),
                    "source_type": source_type,
                }
            )
    return pd.DataFrame(rows, columns=MONETARY_RATES_TNA_REFERENCE_COLUMNS)


def _build_cer_curve() -> tuple[pd.DataFrame, pd.DataFrame]:
    if not CER_BOND_HISTORY_FILE.exists() or not CER_INDEX_FILE.exists():
        empty = pd.DataFrame(columns=CER_ZERO_COUPON_COLUMNS)
        return empty.copy(), empty.copy()
    zero_history, zero_latest = build_cer_zero_coupon_metrics(
        cer_price_history_path=CER_BOND_HISTORY_FILE,
        cer_index_path=CER_INDEX_FILE,
    )
    coupon_history, coupon_latest = build_cer_coupon_metrics(
        cer_price_history_path=CER_BOND_HISTORY_FILE,
        cer_index_path=CER_INDEX_FILE,
    )
    history = pd.concat([zero_history, coupon_history], ignore_index=True)
    latest = pd.concat([zero_latest, coupon_latest], ignore_index=True)
    if history.empty:
        empty = pd.DataFrame(columns=CER_ZERO_COUPON_COLUMNS)
        return empty.copy(), empty.copy()
    history = history[CER_ZERO_COUPON_COLUMNS].sort_values(["date", "payment_date", "symbol"]).reset_index(drop=True)
    latest = latest[CER_ZERO_COUPON_COLUMNS].sort_values(["payment_date", "symbol"]).reset_index(drop=True)
    return history, latest


def _build_hard_dollar_curve() -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    if (
        not IOL_MARKET_PRICES_FILE.exists()
        or not HARD_DOLLAR_REFERENCE_METADATA_FILE.exists()
        or not HARD_DOLLAR_REFERENCE_FLOW_FILE.exists()
    ):
        empty = pd.DataFrame(columns=HARD_DOLLAR_COLUMNS)
        empty_audit = pd.DataFrame()
        return empty.copy(), empty.copy(), empty_audit.copy()
    history, latest, audit = build_hard_dollar_metrics(
        price_history_path=IOL_MARKET_PRICES_FILE,
        reference_metadata_path=HARD_DOLLAR_REFERENCE_METADATA_FILE,
        reference_flow_path=HARD_DOLLAR_REFERENCE_FLOW_FILE,
    )
    return history, latest, audit


def _build_hard_dollar_weighted_parity(hard_dollar_history: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    if hard_dollar_history.empty or not HARD_DOLLAR_REFERENCE_FLOW_FILE.exists():
        empty_history = pd.DataFrame(columns=HARD_DOLLAR_WEIGHTED_PARITY_COLUMNS)
        empty_latest = pd.DataFrame(columns=HARD_DOLLAR_WEIGHTED_PARITY_LATEST_COLUMNS)
        return empty_history.copy(), empty_latest.copy()

    reference_flow = pd.read_csv(HARD_DOLLAR_REFERENCE_FLOW_FILE)
    return build_weighted_parity_history_from_frames(
        history=hard_dollar_history,
        reference_flow=reference_flow,
    )


def _build_fx_futures_curve() -> pd.DataFrame:
    today = pd.Timestamp.today().normalize()

    def keep_live_contracts(frame: pd.DataFrame) -> pd.DataFrame:
        if frame.empty:
            return pd.DataFrame(columns=A3_FUTURES_COLUMNS)
        filtered = frame.copy()
        filtered["maturity_date"] = pd.to_datetime(filtered["maturity_date"], errors="coerce")
        filtered = filtered.dropna(subset=["maturity_date"]).copy()
        filtered = filtered[filtered["maturity_date"] >= today].copy()
        if filtered.empty:
            return pd.DataFrame(columns=A3_FUTURES_COLUMNS)
        return filtered[A3_FUTURES_COLUMNS].sort_values(["maturity_date", "symbol"]).reset_index(drop=True)

    try:
        frame = fetch_a3_dollar_futures_curve()
        if not frame.empty:
            return keep_live_contracts(frame)
    except Exception:
        pass
    if FX_FUTURES_CURVE_LATEST_FILE.exists() and FX_FUTURES_CURVE_LATEST_FILE.stat().st_size > 0:
        cached = pd.read_csv(FX_FUTURES_CURVE_LATEST_FILE, parse_dates=["as_of_date", "maturity_date"])
        return keep_live_contracts(cached)
        return pd.DataFrame(columns=A3_FUTURES_COLUMNS)


def _build_dollar_linked_monitor() -> tuple[pd.DataFrame, pd.DataFrame]:
    empty_history = pd.DataFrame(columns=DOLLAR_LINKED_HISTORY_COLUMNS)
    empty_latest = pd.DataFrame(columns=DOLLAR_LINKED_LATEST_COLUMNS)

    if not IOL_MARKET_PRICES_FILE.exists() or IOL_MARKET_PRICES_FILE.stat().st_size == 0:
        return empty_history.copy(), empty_latest.copy()

    history = pd.read_csv(IOL_MARKET_PRICES_FILE, parse_dates=["date", "fetched_at"])
    history["symbol"] = history["symbol"].astype(str).str.upper()
    history = history[history["symbol"].isin(DOLLAR_LINKED_SYMBOLS)].copy()
    if history.empty:
        return empty_history.copy(), empty_latest.copy()

    history["date"] = pd.to_datetime(history["date"], errors="coerce").dt.normalize()
    history["price"] = pd.to_numeric(history["price"], errors="coerce")
    history["volume"] = pd.to_numeric(history["volume"], errors="coerce")
    history["fetched_at"] = pd.to_datetime(history.get("fetched_at"), errors="coerce")
    history = history.dropna(subset=["date", "symbol"]).copy()
    history = (
        history.sort_values(["date", "symbol", "fetched_at"])
        .drop_duplicates(subset=["date", "symbol"], keep="last")
        .reset_index(drop=True)
    )

    selected_dates = (
        history["date"]
        .dropna()
        .drop_duplicates()
        .sort_values()
        .tolist()[-DOLLAR_LINKED_HISTORY_WINDOW:]
    )
    if not selected_dates:
        return empty_history.copy(), empty_latest.copy()

    history = history[history["date"].isin(selected_dates)].copy()
    history_grid = pd.MultiIndex.from_product(
        [selected_dates, DOLLAR_LINKED_SYMBOLS],
        names=["date", "symbol"],
    ).to_frame(index=False)
    history_filled = history_grid.merge(
        history[["date", "symbol", "price", "volume"]],
        on=["date", "symbol"],
        how="left",
    )
    history_filled["volume"] = history_filled["volume"].fillna(0.0)
    history_filled = history_filled[DOLLAR_LINKED_HISTORY_COLUMNS].sort_values(["date", "symbol"]).reset_index(drop=True)

    history_latest = (
        history.sort_values(["symbol", "date"])
        .groupby("symbol", as_index=False)
        .tail(1)[["date", "symbol", "price", "volume"]]
    )

    if DOLLAR_LINKED_SNAPSHOT_FILE.exists() and DOLLAR_LINKED_SNAPSHOT_FILE.stat().st_size > 0:
        snapshots = pd.read_csv(DOLLAR_LINKED_SNAPSHOT_FILE, parse_dates=["date", "fetched_at"])
        for column in DOLLAR_LINKED_SNAPSHOT_COLUMNS:
            if column not in snapshots.columns:
                snapshots[column] = pd.NA
        snapshots["symbol"] = snapshots["symbol"].astype(str).str.upper()
        snapshots = snapshots[snapshots["symbol"].isin(DOLLAR_LINKED_SYMBOLS)].copy()
        snapshots["date"] = pd.to_datetime(snapshots["date"], errors="coerce").dt.normalize()
        snapshots["price"] = pd.to_numeric(snapshots["price"], errors="coerce")
        snapshots["volume"] = pd.to_numeric(snapshots["volume"], errors="coerce")
        snapshots["amount"] = pd.to_numeric(snapshots["amount"], errors="coerce")
        snapshots["fetched_at"] = pd.to_datetime(snapshots["fetched_at"], errors="coerce")
        snapshot_latest = (
            snapshots.sort_values(["symbol", "date", "fetched_at"])
            .drop_duplicates(subset=["symbol"], keep="last")[["date", "symbol", "price", "volume", "amount"]]
        )
    else:
        snapshot_latest = pd.DataFrame(columns=DOLLAR_LINKED_LATEST_COLUMNS)

    ordered = pd.DataFrame({"symbol": DOLLAR_LINKED_SYMBOLS})
    latest = ordered.merge(history_latest, on="symbol", how="left", suffixes=("", "_history"))
    latest = latest.merge(snapshot_latest, on="symbol", how="left", suffixes=("_history", "_snapshot"))
    latest["date"] = latest["date_snapshot"].combine_first(latest["date_history"])
    latest["price"] = latest["price_snapshot"].combine_first(latest["price_history"])
    latest["volume"] = latest["volume_snapshot"].combine_first(latest["volume_history"])
    latest["amount"] = latest["amount"]
    latest = latest[DOLLAR_LINKED_LATEST_COLUMNS].reset_index(drop=True)

    return history_filled, latest
def _build_monthly_gdp() -> pd.DataFrame:
    empty = pd.DataFrame(columns=["month", "gdp_monthly_ars_mn", "date"])
    if not RAW_GDP_FILE.exists() or RAW_GDP_FILE.stat().st_size == 0:
        return empty

    gdp = pd.read_csv(RAW_GDP_FILE, parse_dates=["quarter_start", "quarter_end"])
    if gdp.empty:
        return empty

    gdp["gdp_current_ars_mn"] = pd.to_numeric(gdp.get("gdp_current_ars_mn"), errors="coerce")
    gdp = gdp.dropna(subset=["quarter_start", "quarter_end", "gdp_current_ars_mn"]).copy()
    rows: list[dict[str, object]] = []
    for row in gdp.itertuples(index=False):
        quarter_start = pd.Timestamp(row.quarter_start).normalize()
        monthly_value = float(row.gdp_current_ars_mn) / 3.0
        for month_offset in range(3):
            month = (quarter_start + pd.DateOffset(months=month_offset)).to_period("M").to_timestamp()
            rows.append(
                {
                    "month": month,
                    "gdp_monthly_ars_mn": monthly_value,
                    "date": month + pd.offsets.MonthEnd(0),
                }
            )
    return pd.DataFrame(rows, columns=["month", "gdp_monthly_ars_mn", "date"]).drop_duplicates(subset=["month"], keep="last")


def _build_monthly_official_fx(prices: pd.DataFrame) -> pd.DataFrame:
    empty = pd.DataFrame(columns=["month", "usd_official", "date"])
    if prices.empty:
        return empty
    frame = prices[prices["asset"] == "usd_ars_official"][["date", "value_nominal"]].copy()
    if frame.empty:
        return empty
    frame["date"] = pd.to_datetime(frame["date"], errors="coerce").dt.normalize()
    frame["value_nominal"] = pd.to_numeric(frame["value_nominal"], errors="coerce")
    frame = frame.dropna(subset=["date", "value_nominal"]).copy()
    frame["month"] = frame["date"].dt.to_period("M").dt.to_timestamp()
    monthly = (
        frame.sort_values(["month", "date"])
        .groupby("month", as_index=False)
        .tail(1)
        .rename(columns={"value_nominal": "usd_official"})
        .reset_index(drop=True)
    )
    return monthly[["month", "usd_official", "date"]]


def _build_monetary_variables(
    prices: pd.DataFrame,
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    empty_latest = pd.DataFrame(columns=MONETARY_VARIABLE_CORE_LATEST_COLUMNS)
    empty_economy = pd.DataFrame(columns=MONETARY_ECONOMY_PIB_HISTORY_COLUMNS)
    empty_peso_usd = pd.DataFrame(columns=MONETARY_PESO_USD_PIB_HISTORY_COLUMNS)
    empty_usd_stocks = pd.DataFrame(columns=MONETARY_USD_STOCKS_HISTORY_COLUMNS)
    empty_freshness = pd.DataFrame(columns=MONETARY_BLOCK_FRESHNESS_COLUMNS)
    empty_source_table = pd.DataFrame(columns=MONETARY_SOURCE_TABLE_COLUMNS)

    if not MONETARY_VARIABLES_RAW_FILE.exists() or MONETARY_VARIABLES_RAW_FILE.stat().st_size == 0:
        return (
            empty_latest.copy(),
            empty_economy.copy(),
            empty_peso_usd.copy(),
            empty_usd_stocks.copy(),
            empty_freshness.copy(),
            empty_source_table.copy(),
        )

    raw = pd.read_csv(MONETARY_VARIABLES_RAW_FILE, parse_dates=["date"])
    if raw.empty:
        return (
            empty_latest.copy(),
            empty_economy.copy(),
            empty_peso_usd.copy(),
            empty_usd_stocks.copy(),
            empty_freshness.copy(),
            empty_source_table.copy(),
        )

    raw["metric"] = raw.get("metric", pd.Series(dtype="object")).astype(str)
    raw["date"] = pd.to_datetime(raw.get("date"), errors="coerce").dt.normalize()
    raw["value"] = pd.to_numeric(raw.get("value"), errors="coerce")
    raw["source_id"] = pd.to_numeric(raw.get("source_id"), errors="coerce")
    raw["label"] = raw["metric"].map(MONETARY_VARIABLE_LABELS).fillna(raw.get("label"))
    raw = raw[raw["metric"].isin(MONETARY_VARIABLE_METRICS)].dropna(subset=["date", "metric", "value"]).copy()
    raw = (
        raw.sort_values(["metric", "date", "source_id"])
        .drop_duplicates(subset=["metric", "date"], keep="last")
        .reset_index(drop=True)
    )
    if raw.empty:
        return (
            empty_latest.copy(),
            empty_economy.copy(),
            empty_peso_usd.copy(),
            empty_usd_stocks.copy(),
            empty_freshness.copy(),
            empty_source_table.copy(),
        )

    reference_dates = [pd.Timestamp(raw["date"].max()).normalize()]
    if not prices.empty and "date" in prices.columns:
        price_dates = pd.to_datetime(prices["date"], errors="coerce").dropna()
        if not price_dates.empty:
            reference_dates.append(pd.Timestamp(price_dates.max()).normalize())
    overall_latest_date = max(reference_dates)

    latest_rows: list[dict[str, object]] = []
    source_rows: list[dict[str, object]] = []
    for spec in MONETARY_VARIABLE_SPECS:
        metric_frame = raw[raw["metric"] == spec["metric"]].sort_values("date").reset_index(drop=True)
        if metric_frame.empty:
            continue

        latest_row = metric_frame.iloc[-1]
        latest_date = pd.Timestamp(latest_row["date"]).normalize()

        def baseline_pct(day_offset: int) -> float | None:
            baseline = metric_frame[metric_frame["date"] <= latest_date - pd.Timedelta(days=day_offset)]
            if baseline.empty:
                return None
            base_value = float(baseline.iloc[-1]["value"])
            if base_value == 0:
                return None
            return (float(latest_row["value"]) / base_value - 1.0) * 100.0

        if spec["metric"] in MONETARY_CORE_METRICS:
            latest_rows.append(
                {
                    "metric": spec["metric"],
                    "label": spec["label"],
                    "date": latest_date,
                    "value": float(latest_row["value"]),
                    "unit": spec["unit"],
                    "source_id": int(spec["source_id"]),
                    "change_30d_pct": baseline_pct(30),
                    "change_yoy_pct": baseline_pct(365),
                    "lag_days": int((overall_latest_date - latest_date).days),
                }
            )

        source_rows.append(
            {
                "series_key": spec["metric"],
                "label": spec["label"],
                "date": latest_date,
                "value": float(latest_row["value"]),
                "unit": spec["unit"],
                "lag_days": int((overall_latest_date - latest_date).days),
            }
        )

    latest = pd.DataFrame(latest_rows, columns=MONETARY_VARIABLE_CORE_LATEST_COLUMNS)
    latest["metric_rank"] = latest["metric"].map({metric: idx for idx, metric in enumerate(MONETARY_CORE_METRICS)})
    latest = latest.sort_values("metric_rank").drop(columns=["metric_rank"]).reset_index(drop=True)

    monthly = raw.copy()
    monthly["month"] = monthly["date"].dt.to_period("M").dt.to_timestamp()
    monthly = (
        monthly.sort_values(["metric", "month", "date"])
        .groupby(["metric", "month"], as_index=False)
        .tail(1)
        .reset_index(drop=True)
    )

    gdp_monthly = _build_monthly_gdp()
    if not gdp_monthly.empty:
        source_rows.append(
            {
                "series_key": "gdp_nominal_monthly",
                "label": "PIB nominal mensualizado",
                "date": pd.Timestamp(gdp_monthly["date"].max()).normalize(),
                "value": float(gdp_monthly.sort_values("date").iloc[-1]["gdp_monthly_ars_mn"]),
                "unit": MONETARY_VARIABLE_UNIT,
                "lag_days": int((overall_latest_date - pd.Timestamp(gdp_monthly["date"].max()).normalize()).days),
            }
        )

    usd_monthly = _build_monthly_official_fx(prices)
    if not usd_monthly.empty:
        source_rows.append(
            {
                "series_key": "usd_ars_official",
                "label": "USD oficial A3500",
                "date": pd.Timestamp(usd_monthly["date"].max()).normalize(),
                "value": float(usd_monthly.sort_values("date").iloc[-1]["usd_official"]),
                "unit": "ARS/USD",
                "lag_days": int((overall_latest_date - pd.Timestamp(usd_monthly["date"].max()).normalize()).days),
            }
        )

    monthly_pivot = (
        monthly.pivot_table(index="month", columns="metric", values="value", aggfunc="last")
        .sort_index()
        .reset_index()
    )

    economy = empty_economy.copy()
    if not monthly_pivot.empty and not gdp_monthly.empty:
        economy = monthly_pivot.merge(gdp_monthly[["month", "gdp_monthly_ars_mn"]], on="month", how="inner")
        required = [
            "circulacion_monetaria",
            "depositos_vista_privado",
            "depositos_plazo_privado",
            "m2_total",
            "m3_total",
            "gdp_monthly_ars_mn",
        ]
        economy = economy.dropna(subset=required).copy()
        if not economy.empty:
            economy["month_label"] = economy["month"].dt.strftime("%b-%y").str.lower()
            economy["circulacion_pct_pib"] = economy["circulacion_monetaria"] / economy["gdp_monthly_ars_mn"] * 100.0
            economy["depositos_vista_pct_pib"] = economy["depositos_vista_privado"] / economy["gdp_monthly_ars_mn"] * 100.0
            economy["depositos_plazo_pct_pib"] = economy["depositos_plazo_privado"] / economy["gdp_monthly_ars_mn"] * 100.0
            economy["m2_pct_pib"] = economy["m2_total"] / economy["gdp_monthly_ars_mn"] * 100.0
            economy["m3_pct_pib"] = economy["m3_total"] / economy["gdp_monthly_ars_mn"] * 100.0
            selected_months = economy["month"].drop_duplicates().sort_values().tolist()[-MONETARY_VARIABLE_WINDOW_MONTHS:]
            economy = economy[economy["month"].isin(selected_months)].copy()
            economy = economy[MONETARY_ECONOMY_PIB_HISTORY_COLUMNS].reset_index(drop=True)

    peso_usd = empty_peso_usd.copy()
    if not monthly_pivot.empty and not gdp_monthly.empty and not usd_monthly.empty:
        peso_usd = (
            monthly_pivot
            .merge(gdp_monthly[["month", "gdp_monthly_ars_mn"]], on="month", how="inner")
            .merge(usd_monthly[["month", "usd_official"]], on="month", how="inner")
        )
        required = [
            "circulacion_monetaria",
            "depositos_privados_ars",
            "depositos_privados_usd",
            "gdp_monthly_ars_mn",
            "usd_official",
        ]
        peso_usd = peso_usd.dropna(subset=required).copy()
        if not peso_usd.empty:
            peso_usd["month_label"] = peso_usd["month"].dt.strftime("%b-%y").str.lower()
            peso_usd["depositos_usd_ars"] = peso_usd["depositos_privados_usd"] * peso_usd["usd_official"]
            peso_usd["circulacion_pct_pib"] = peso_usd["circulacion_monetaria"] / peso_usd["gdp_monthly_ars_mn"] * 100.0
            peso_usd["depositos_pesos_pct_pib"] = peso_usd["depositos_privados_ars"] / peso_usd["gdp_monthly_ars_mn"] * 100.0
            peso_usd["depositos_usd_pct_pib"] = peso_usd["depositos_usd_ars"] / peso_usd["gdp_monthly_ars_mn"] * 100.0
            peso_usd["total_pct_pib"] = (
                peso_usd["circulacion_pct_pib"] + peso_usd["depositos_pesos_pct_pib"] + peso_usd["depositos_usd_pct_pib"]
            )
            selected_months = peso_usd["month"].drop_duplicates().sort_values().tolist()[-MONETARY_VARIABLE_WINDOW_MONTHS:]
            peso_usd = peso_usd[peso_usd["month"].isin(selected_months)].copy()
            peso_usd = peso_usd[MONETARY_PESO_USD_PIB_HISTORY_COLUMNS].reset_index(drop=True)

    usd_stocks = empty_usd_stocks.copy()
    if not monthly_pivot.empty and not usd_monthly.empty:
        usd_stocks = monthly_pivot.merge(usd_monthly[["month", "usd_official"]], on="month", how="inner")
        required = [
            "circulacion_monetaria",
            "m2_total",
            "depositos_privados_ars",
            "depositos_privados_usd",
            "usd_official",
        ]
        usd_stocks = usd_stocks.dropna(subset=required).copy()
        if not usd_stocks.empty:
            usd_stocks["month_label"] = usd_stocks["month"].dt.strftime("%b-%y").str.lower()
            usd_stocks["m2_usd_mn"] = usd_stocks["m2_total"] / usd_stocks["usd_official"]
            usd_stocks["depositos_pesos_usd_mn"] = usd_stocks["depositos_privados_ars"] / usd_stocks["usd_official"]
            usd_stocks["depositos_usd_mn"] = usd_stocks["depositos_privados_usd"]
            usd_stocks["total_usd_mn"] = (
                usd_stocks["circulacion_monetaria"] / usd_stocks["usd_official"]
                + usd_stocks["depositos_pesos_usd_mn"]
                + usd_stocks["depositos_usd_mn"]
            )
            selected_months = usd_stocks["month"].drop_duplicates().sort_values().tolist()[-MONETARY_VARIABLE_WINDOW_MONTHS:]
            usd_stocks = usd_stocks[usd_stocks["month"].isin(selected_months)].copy()
            usd_stocks = usd_stocks[MONETARY_USD_STOCKS_HISTORY_COLUMNS].reset_index(drop=True)

    freshness_rows: list[dict[str, object]] = []
    for block_key, label, frame, date_column in [
        ("core_latest", "Core diario", latest, "date"),
        ("economy_pib_history", "Monetizacion de la economia", economy, "month"),
        ("peso_usd_pib_history", "Monetizacion pesos y dolares", peso_usd, "month"),
        ("usd_stocks_history", "Stocks monetarios en USD", usd_stocks, "month"),
    ]:
        parsed = pd.to_datetime(frame.get(date_column, pd.Series(dtype="datetime64[ns]")), errors="coerce").dropna()
        latest_date = pd.Timestamp(parsed.max()).normalize() if not parsed.empty else pd.NaT
        lag_days = int((overall_latest_date - latest_date).days) if not pd.isna(latest_date) else pd.NA
        freshness_rows.append(
            {
                "block": block_key,
                "label": label,
                "latest_date": latest_date,
                "window_days": MONETARY_BLOCK_WINDOWS.get(block_key, 3),
                "lag_days": lag_days,
            }
        )
    block_freshness = pd.DataFrame(freshness_rows, columns=MONETARY_BLOCK_FRESHNESS_COLUMNS)
    source_table = pd.DataFrame(source_rows, columns=MONETARY_SOURCE_TABLE_COLUMNS).sort_values(["date", "label"]).reset_index(drop=True)

    return latest, economy, peso_usd, usd_stocks, block_freshness, source_table


def _build_monetary_rates_tna(
    observed_inflation: pd.DataFrame,
    bundle_reference_date: pd.Timestamp | None = None,
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    empty_series = pd.DataFrame(columns=MONETARY_RATES_TNA_SERIES_COLUMNS)
    empty_reference = pd.DataFrame(columns=MONETARY_RATES_TNA_REFERENCE_COLUMNS)
    empty_latest = pd.DataFrame(columns=MONETARY_RATES_TNA_LATEST_COLUMNS)
    empty_freshness = pd.DataFrame(columns=MONETARY_BLOCK_FRESHNESS_COLUMNS)

    if not MONETARY_RATES_TNA_RAW_FILE.exists() or MONETARY_RATES_TNA_RAW_FILE.stat().st_size == 0:
        return empty_series.copy(), empty_reference.copy(), empty_latest.copy(), empty_freshness.copy()

    raw = pd.read_csv(MONETARY_RATES_TNA_RAW_FILE, parse_dates=["date"])
    if raw.empty:
        return empty_series.copy(), empty_reference.copy(), empty_latest.copy(), empty_freshness.copy()

    raw["metric"] = raw.get("metric", pd.Series(dtype="object")).astype(str)
    raw["date"] = pd.to_datetime(raw.get("date"), errors="coerce").dt.normalize()
    raw["value"] = pd.to_numeric(raw.get("value"), errors="coerce")
    raw["source_id"] = pd.to_numeric(raw.get("source_id"), errors="coerce")
    raw["label"] = raw["metric"].map(MONETARY_RATES_TNA_LABELS).fillna(raw.get("label"))
    raw = raw[raw["metric"].isin(MONETARY_RATES_TNA_METRICS)].dropna(subset=["date", "metric", "value"]).copy()
    raw = (
        raw.sort_values(["metric", "date", "source_id"])
        .drop_duplicates(subset=["metric", "date"], keep="last")
        .reset_index(drop=True)
    )
    if raw.empty:
        return empty_series.copy(), empty_reference.copy(), empty_latest.copy(), empty_freshness.copy()

    latest_series_date = pd.Timestamp(raw["date"].max()).normalize()
    if bundle_reference_date is not None and not pd.isna(bundle_reference_date):
        reference_anchor = max(latest_series_date, pd.Timestamp(bundle_reference_date).normalize())
    else:
        reference_anchor = latest_series_date
    window_start = (latest_series_date - pd.DateOffset(months=MONETARY_VARIABLE_WINDOW_MONTHS)).normalize()

    series = raw[raw["date"] >= window_start][["date", "metric", "label", "value"]].copy()
    series = series.sort_values(["metric", "date"]).reset_index(drop=True)

    observed = observed_inflation[["month", "inflation_monthly"]].copy()
    observed["month"] = pd.to_datetime(observed["month"], errors="coerce").dt.to_period("M").dt.to_timestamp()
    observed = observed.dropna(subset=["month", "inflation_monthly"]).copy()
    observed["value"] = _annualize_monthly_inflation(observed["inflation_monthly"])
    observed_daily = _expand_monthly_reference_daily(
        observed[["month", "value"]],
        metric="inflacion_observada_anualizada",
        label="Inflacion observada anualizada",
        source_type="observed",
    )

    ecogo = _read_ecogo_inflation_projection()
    if not observed.empty:
        first_forecast_month = observed["month"].max() + pd.offsets.MonthBegin(1)
        ecogo = ecogo[ecogo["month"] >= first_forecast_month].copy()
    ecogo["value"] = _annualize_monthly_inflation(ecogo.get("inflation_monthly"))
    forecast_daily = _expand_monthly_reference_daily(
        ecogo[["month", "value"]],
        metric="inflacion_esperada_ecogo_anualizada",
        label="Inflacion esperada EcoGo anualizada",
        source_type="forecast",
    )

    reference = pd.concat([observed_daily, forecast_daily], ignore_index=True)
    reference["date"] = pd.to_datetime(reference["date"], errors="coerce").dt.normalize()
    reference["value"] = pd.to_numeric(reference["value"], errors="coerce")
    reference = (
        reference.dropna(subset=["date", "metric", "value"])
        .sort_values(["metric", "date"])
        .drop_duplicates(subset=["metric", "date"], keep="last")
        .reset_index(drop=True)
    )

    if not reference.empty:
        reference = reference[reference["date"] >= window_start].copy()

    reference_join = (
        reference[["date", "metric", "value"]]
        .rename(columns={"metric": "reference_metric", "value": "reference_value"})
        .sort_values(["reference_metric", "date"])
        .reset_index(drop=True)
    )
    observed_reference_end = observed["month"].max() + pd.offsets.MonthEnd(0) if not observed.empty else pd.NaT
    latest_rows: list[dict[str, object]] = []
    for spec in MONETARY_RATES_TNA_SPECS:
        metric_frame = raw[raw["metric"] == spec["metric"]].sort_values("date").reset_index(drop=True)
        if metric_frame.empty:
            continue
        latest_row = metric_frame.iloc[-1]
        latest_date = pd.Timestamp(latest_row["date"]).normalize()
        reference_metric = "inflacion_observada_anualizada"
        if pd.isna(observed_reference_end) or latest_date > observed_reference_end:
            reference_metric = "inflacion_esperada_ecogo_anualizada"
        ref_frame = reference_join[reference_join["reference_metric"] == reference_metric].copy()
        ref_value = None
        if not ref_frame.empty:
            ref_match = ref_frame[ref_frame["date"] == latest_date]
            if ref_match.empty:
                ref_match = ref_frame[ref_frame["date"] <= latest_date].tail(1)
            if not ref_match.empty:
                ref_value = float(ref_match.iloc[-1]["reference_value"])
        latest_rows.append(
            {
                "metric": spec["metric"],
                "label": spec["label"],
                "date": latest_date,
                "value": float(latest_row["value"]),
                "lag_days": int((reference_anchor - latest_date).days),
                "spread_vs_reference_pct": float(latest_row["value"]) - ref_value if ref_value is not None else pd.NA,
            }
        )

    latest = pd.DataFrame(latest_rows, columns=MONETARY_RATES_TNA_LATEST_COLUMNS)

    block_freshness = pd.DataFrame(
        [
            {
                "block": "rates_tna",
                "label": "Tasas TNA",
                "latest_date": latest_series_date,
                "window_days": MONETARY_BLOCK_WINDOWS.get("rates_tna", 5),
                "lag_days": int((reference_anchor - latest_series_date).days),
            }
        ],
        columns=MONETARY_BLOCK_FRESHNESS_COLUMNS,
    )

    series = series[MONETARY_RATES_TNA_SERIES_COLUMNS].reset_index(drop=True)
    reference = reference[MONETARY_RATES_TNA_REFERENCE_COLUMNS].reset_index(drop=True)
    latest = latest.sort_values("date").reset_index(drop=True)
    return series, reference, latest, block_freshness


def _build_fci_category_history() -> pd.DataFrame:
    existing = None
    if FCI_CATEGORY_HISTORY_FILE.exists() and FCI_CATEGORY_HISTORY_FILE.stat().st_size > 0:
        existing = pd.read_csv(FCI_CATEGORY_HISTORY_FILE, parse_dates=["date"])
    try:
        frame = fetch_fci_category_history(
            existing=existing,
            lookback_days=FCI_FETCH_LOOKBACK_DAYS,
            max_workers=FCI_FETCH_MAX_WORKERS,
            timeout=FCI_FETCH_TIMEOUT_SECONDS,
        )
        if not frame.empty:
            return frame
    except Exception:
        pass
    if existing is not None:
        return existing
    return pd.DataFrame(columns=FCI_CATEGORY_COLUMNS)


def _build_fci_currency_history() -> pd.DataFrame:
    existing = None
    if FCI_CURRENCY_HISTORY_FILE.exists() and FCI_CURRENCY_HISTORY_FILE.stat().st_size > 0:
        existing = pd.read_csv(FCI_CURRENCY_HISTORY_FILE, parse_dates=["date"])
    try:
        frame = fetch_fci_currency_history(
            existing=existing,
            lookback_days=FCI_FETCH_LOOKBACK_DAYS,
            max_workers=FCI_FETCH_MAX_WORKERS,
            timeout=FCI_FETCH_TIMEOUT_SECONDS,
        )
        if not frame.empty:
            return frame
    except Exception:
        pass
    if existing is not None:
        return existing
    return pd.DataFrame(columns=FCI_CURRENCY_COLUMNS)


def _build_fci_horizon_history() -> pd.DataFrame:
    existing = None
    if FCI_HORIZON_HISTORY_FILE.exists() and FCI_HORIZON_HISTORY_FILE.stat().st_size > 0:
        existing = pd.read_csv(FCI_HORIZON_HISTORY_FILE, parse_dates=["date"])
    try:
        frame = fetch_fci_horizon_history(
            existing=existing,
            lookback_days=FCI_FETCH_LOOKBACK_DAYS,
            max_workers=FCI_FETCH_MAX_WORKERS,
            timeout=FCI_FETCH_TIMEOUT_SECONDS,
        )
        if not frame.empty:
            return frame
    except Exception:
        pass
    if existing is not None:
        return existing
    return pd.DataFrame(columns=FCI_HORIZON_COLUMNS)


def _build_fci_money_market_latest() -> pd.DataFrame:
    try:
        frame = fetch_money_market_fund_latest(timeout=FCI_FETCH_TIMEOUT_SECONDS)
        if not frame.empty:
            return frame
    except Exception:
        pass
    if FCI_MONEY_MARKET_LATEST_FILE.exists() and FCI_MONEY_MARKET_LATEST_FILE.stat().st_size > 0:
        return pd.read_csv(FCI_MONEY_MARKET_LATEST_FILE, parse_dates=["date"])
    return pd.DataFrame(columns=FCI_MONEY_MARKET_COLUMNS)


def _build_fci_reported_slice_history() -> pd.DataFrame:
    existing = None
    if FCI_REPORTED_SLICE_HISTORY_FILE.exists() and FCI_REPORTED_SLICE_HISTORY_FILE.stat().st_size > 0:
        existing = pd.read_csv(FCI_REPORTED_SLICE_HISTORY_FILE, parse_dates=["date"])
    try:
        frame = fetch_fci_reported_slice_history(
            existing=existing,
            lookback_days=FCI_FETCH_LOOKBACK_DAYS,
            max_workers=FCI_FETCH_MAX_WORKERS,
            timeout=FCI_FETCH_TIMEOUT_SECONDS,
        )
        if not frame.empty:
            return frame
    except Exception:
        pass
    if existing is not None:
        return existing
    return pd.DataFrame(columns=FCI_REPORTED_SLICE_COLUMNS)


def _resolve_iol_metadata_workbook() -> Path:
    if DEFAULT_IOL_INSTRUMENTS_DROP_PATH.exists():
        return DEFAULT_IOL_INSTRUMENTS_DROP_PATH
    return DEFAULT_IOL_INSTRUMENTS_METADATA_PATH


def _build_lecap_macro_history(
    prices: pd.DataFrame,
    observed_inflation: pd.DataFrame,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    empty_history = pd.DataFrame(columns=LECAP_MACRO_HISTORY_COLUMNS)
    empty_monthly = pd.DataFrame(columns=LECAP_MACRO_MONTHLY_COLUMNS)

    if not IOL_ZERO_COUPON_METRICS_FILE.exists() or IOL_ZERO_COUPON_METRICS_FILE.stat().st_size == 0:
        return empty_history.copy(), empty_monthly.copy()

    metadata_path = _resolve_iol_metadata_workbook()
    if not metadata_path.exists():
        return empty_history.copy(), empty_monthly.copy()

    metadata = load_iol_instrument_metadata(metadata_path)
    lecap_symbols = (
        metadata[metadata["instrument_family"].astype(str).str.lower() == "lecap"]["symbol"]
        .dropna()
        .astype(str)
        .str.upper()
        .unique()
        .tolist()
    )
    if not lecap_symbols:
        return empty_history.copy(), empty_monthly.copy()

    metrics = pd.read_csv(
        IOL_ZERO_COUPON_METRICS_FILE,
        usecols=["date", "symbol", "instrument_family", "price", "tem", "tea"],
        parse_dates=["date"],
    )
    if metrics.empty:
        return empty_history.copy(), empty_monthly.copy()

    metrics["symbol"] = metrics["symbol"].astype(str).str.upper()
    metrics["instrument_family"] = metrics["instrument_family"].astype(str).str.lower()
    history = metrics[
        (metrics["instrument_family"] == "lecap")
        & (metrics["symbol"].isin(lecap_symbols))
    ].copy()
    if history.empty:
        return empty_history.copy(), empty_monthly.copy()

    history["price"] = pd.to_numeric(history["price"], errors="coerce")
    history["tem_pct"] = pd.to_numeric(history["tem"], errors="coerce") * 100.0
    history["tea_pct"] = pd.to_numeric(history["tea"], errors="coerce") * 100.0
    history = history.dropna(subset=["date", "symbol", "tem_pct"]).copy()
    if history.empty:
        return empty_history.copy(), empty_monthly.copy()

    # Exclude the final two observations per LECAP so the chart stops at t-2
    # and never uses the maturity print or the immediately preceding datum.
    history = history.sort_values(["symbol", "date"]).reset_index(drop=True)
    history["obs_rank"] = history.groupby("symbol").cumcount()
    history["obs_count"] = history.groupby("symbol")["symbol"].transform("size")
    history["obs_from_end"] = history["obs_count"] - history["obs_rank"] - 1
    history = history[history["obs_from_end"] >= 2].copy()
    if history.empty:
        return empty_history.copy(), empty_monthly.copy()
    history = history.drop(columns=["obs_rank", "obs_count", "obs_from_end"])

    averages = (
        history.groupby("date", as_index=False)
        .agg(
            average_tem_pct=("tem_pct", "mean"),
            average_price=("price", "mean"),
            symbol_count=("symbol", "nunique"),
        )
        .sort_values("date")
        .reset_index(drop=True)
    )
    history = (
        history.merge(averages, on="date", how="left")
        [LECAP_MACRO_HISTORY_COLUMNS]
        .sort_values(["date", "symbol"])
        .reset_index(drop=True)
    )

    first_month = history["date"].min().to_period("M").to_timestamp()
    last_month = history["date"].max().to_period("M").to_timestamp()

    inflation = observed_inflation[["month", "month_label", "inflation_monthly_pct"]].copy()
    a3500 = prices[prices["asset"] == "usd_ars_official"][["date", "value_nominal"]].copy()
    if a3500.empty:
        monthly = inflation.copy()
        monthly["a3500_avg"] = pd.NA
        monthly["a3500_devaluation_pct"] = pd.NA
    else:
        a3500["month"] = a3500["date"].dt.to_period("M").dt.to_timestamp()
        a3500_monthly = (
            a3500.groupby("month", as_index=False)
            .agg(a3500_avg=("value_nominal", "mean"))
            .sort_values("month")
            .reset_index(drop=True)
        )
        a3500_monthly["a3500_devaluation_pct"] = a3500_monthly["a3500_avg"].pct_change() * 100.0
        monthly = inflation.merge(a3500_monthly, on="month", how="outer").sort_values("month").reset_index(drop=True)
        monthly["month_label"] = monthly["month_label"].fillna(monthly["month"].dt.strftime("%Y-%m"))

    monthly = monthly[
        (monthly["month"] >= first_month)
        & (monthly["month"] <= last_month)
    ].copy()
    return history, monthly[LECAP_MACRO_MONTHLY_COLUMNS].reset_index(drop=True)


def run() -> dict[str, Path]:
    ensure_data_dirs()

    prices = pd.read_csv(RAW_PRICES_FILE, parse_dates=["date"])
    cpi = pd.read_csv(RAW_CPI_FILE, parse_dates=["date"])
    fundamentals = pd.read_csv(RAW_FUNDAMENTALS_FILE)

    prices["month"] = prices["date"].dt.to_period("M").dt.to_timestamp()
    prices = prices.sort_values("month").reset_index(drop=True)
    cpi = cpi.rename(columns={"date": "month"}).sort_values("month").reset_index(drop=True)
    latest_cpi = float(cpi["cpi_index"].iloc[-1])

    merged = pd.merge_asof(
        prices,
        cpi,
        on="month",
        direction="backward",
        allow_exact_matches=True,
    )
    merged["value_real_today"] = merged["value_nominal"] * latest_cpi / merged["cpi_index"]
    merged = _apply_performance_basis(merged)
    series = (
        pd.concat([_compute_group_metrics(group) for _, group in merged.groupby("asset", sort=False)], ignore_index=True)
        .sort_values(["category", "asset_label", "date"])
        .reset_index(drop=True)
    )
    series = series.drop(columns=["month"])

    latest = (
        series.sort_values("date")
        .groupby("asset", as_index=False)
        .tail(1)
        .sort_values(["category", "asset_label"])
        .reset_index(drop=True)
    )

    fundamentals = fundamentals.sort_values(["category", "asset_label", "year"]).reset_index(drop=True)
    fundamentals_latest = (
        fundamentals.sort_values("year")
        .groupby(["asset", "scope", "metric"], as_index=False)
        .tail(1)
        .sort_values(["category", "asset_label", "metric_label"])
        .reset_index(drop=True)
    )

    alerts = _build_alerts(latest)
    observed_inflation = _build_observed_inflation(cpi)
    cer_curve_history, cer_curve_latest = _build_cer_curve()
    hard_dollar_history, hard_dollar_latest, hard_dollar_audit = _build_hard_dollar_curve()
    hard_dollar_weighted_parity_history, hard_dollar_weighted_parity_latest = _build_hard_dollar_weighted_parity(
        hard_dollar_history
    )
    fx_futures_curve_latest = _build_fx_futures_curve()
    dollar_linked_history, dollar_linked_latest = _build_dollar_linked_monitor()
    (
        monetary_variables_latest,
        monetary_variables_economy_pib_history,
        monetary_variables_peso_usd_pib_history,
        monetary_variables_usd_stocks_history,
        monetary_variables_block_freshness,
        monetary_variables_source_table,
    ) = _build_monetary_variables(prices)
    (
        monetary_rates_tna_series,
        monetary_rates_tna_reference,
        monetary_rates_tna_latest,
        monetary_rates_tna_freshness,
    ) = _build_monetary_rates_tna(
        observed_inflation,
        bundle_reference_date=pd.to_datetime(prices["date"], errors="coerce").dropna().max() if not prices.empty else None,
    )
    monetary_variables_block_freshness = (
        pd.concat([monetary_variables_block_freshness, monetary_rates_tna_freshness], ignore_index=True)
        if not monetary_rates_tna_freshness.empty
        else monetary_variables_block_freshness
    )
    if not monetary_variables_block_freshness.empty:
        monetary_variables_block_freshness = (
            monetary_variables_block_freshness
            .sort_values(["block", "latest_date"])
            .drop_duplicates(subset=["block"], keep="last")
            .reset_index(drop=True)
        )
    fci_category_history = _build_fci_category_history()
    fci_currency_history = _build_fci_currency_history()
    fci_horizon_history = _build_fci_horizon_history()
    fci_money_market_latest = _build_fci_money_market_latest()
    fci_reported_slice_history = _build_fci_reported_slice_history()
    lecap_history, lecap_latest = _build_lecap_monitor()
    lecap_macro_history, lecap_macro_monthly = _build_lecap_macro_history(prices, observed_inflation)
    breakeven_history, breakeven_latest = build_breakeven_monitor_live(
        fixed_metrics=lecap_history,
        cer_price_history=cer_curve_history[["date", "symbol", "price"]].copy(),
        cer_index_history=load_cer_index_history(CER_INDEX_FILE),
    )
    breakeven_monthly_history, breakeven_monthly_latest = build_breakeven_forward_monthly(breakeven_history)

    series.to_csv(SERIES_FILE, index=False)
    latest.to_csv(LATEST_FILE, index=False)
    fundamentals.to_csv(FUNDAMENTALS_FILE, index=False)
    fundamentals_latest.to_csv(FUNDAMENTALS_LATEST_FILE, index=False)
    alerts.to_csv(ALERTS_FILE, index=False)
    lecap_history.to_csv(LECAP_MONITOR_HISTORY_FILE, index=False)
    lecap_latest.to_csv(LECAP_MONITOR_LATEST_FILE, index=False)
    breakeven_history.to_csv(BREAKEVEN_HISTORY_FILE, index=False)
    breakeven_latest.to_csv(BREAKEVEN_LATEST_FILE, index=False)
    breakeven_monthly_history.to_csv(BREAKEVEN_MONTHLY_HISTORY_FILE, index=False)
    breakeven_monthly_latest.to_csv(BREAKEVEN_MONTHLY_LATEST_FILE, index=False)
    observed_inflation.to_csv(OBSERVED_INFLATION_FILE, index=False)
    lecap_macro_history.to_csv(LECAP_MACRO_HISTORY_FILE, index=False)
    lecap_macro_monthly.to_csv(LECAP_MACRO_MONTHLY_FILE, index=False)
    cer_curve_history.to_csv(CER_CURVE_HISTORY_FILE, index=False)
    cer_curve_latest.to_csv(CER_CURVE_LATEST_FILE, index=False)
    hard_dollar_history.to_csv(HARD_DOLLAR_CURVE_HISTORY_FILE, index=False)
    hard_dollar_latest.to_csv(HARD_DOLLAR_CURVE_LATEST_FILE, index=False)
    hard_dollar_audit.to_csv(HARD_DOLLAR_AUDIT_FILE, index=False)
    hard_dollar_weighted_parity_history.to_csv(HARD_DOLLAR_WEIGHTED_PARITY_HISTORY_FILE, index=False)
    hard_dollar_weighted_parity_latest.to_csv(HARD_DOLLAR_WEIGHTED_PARITY_LATEST_FILE, index=False)
    fx_futures_curve_latest.to_csv(FX_FUTURES_CURVE_LATEST_FILE, index=False)
    dollar_linked_history.to_csv(DOLLAR_LINKED_HISTORY_FILE, index=False)
    dollar_linked_latest.to_csv(DOLLAR_LINKED_LATEST_FILE, index=False)
    monetary_variables_latest.to_csv(MONETARY_VARIABLES_LATEST_FILE, index=False)
    monetary_variables_economy_pib_history.to_csv(MONETARY_VARIABLES_ECONOMY_PIB_HISTORY_FILE, index=False)
    monetary_variables_peso_usd_pib_history.to_csv(MONETARY_VARIABLES_PESO_USD_PIB_HISTORY_FILE, index=False)
    monetary_variables_usd_stocks_history.to_csv(MONETARY_VARIABLES_USD_STOCKS_HISTORY_FILE, index=False)
    monetary_variables_block_freshness.to_csv(MONETARY_VARIABLES_BLOCK_FRESHNESS_FILE, index=False)
    monetary_variables_source_table.to_csv(MONETARY_VARIABLES_SOURCE_TABLE_FILE, index=False)
    monetary_rates_tna_series.to_csv(MONETARY_RATES_TNA_SERIES_FILE, index=False)
    monetary_rates_tna_reference.to_csv(MONETARY_RATES_TNA_REFERENCE_FILE, index=False)
    monetary_rates_tna_latest.to_csv(MONETARY_RATES_TNA_LATEST_FILE, index=False)
    fci_category_history.to_csv(FCI_CATEGORY_HISTORY_FILE, index=False)
    fci_currency_history.to_csv(FCI_CURRENCY_HISTORY_FILE, index=False)
    fci_horizon_history.to_csv(FCI_HORIZON_HISTORY_FILE, index=False)
    fci_money_market_latest.to_csv(FCI_MONEY_MARKET_LATEST_FILE, index=False)
    fci_reported_slice_history.to_csv(FCI_REPORTED_SLICE_HISTORY_FILE, index=False)

    metadata = {
        "generated_at": pd.Timestamp.utcnow().isoformat(),
        "assets": int(latest["asset"].nunique()),
        "categories": sorted(latest["category"].dropna().unique().tolist()),
        "start_date": series["date"].min().date().isoformat(),
        "end_date": series["date"].max().date().isoformat(),
        "note": _metadata_note(),
    }
    META_FILE.write_text(json.dumps(metadata, indent=2), encoding="utf-8")

    return {
        "prices": SERIES_FILE,
        "latest_prices": LATEST_FILE,
        "fundamentals": FUNDAMENTALS_FILE,
        "fundamentals_latest": FUNDAMENTALS_LATEST_FILE,
        "alerts": ALERTS_FILE,
        "lecap_monitor_history": LECAP_MONITOR_HISTORY_FILE,
        "lecap_monitor_latest": LECAP_MONITOR_LATEST_FILE,
        "cer_breakeven_history": BREAKEVEN_HISTORY_FILE,
        "cer_breakeven_latest": BREAKEVEN_LATEST_FILE,
        "cer_breakeven_monthly_history": BREAKEVEN_MONTHLY_HISTORY_FILE,
        "cer_breakeven_monthly_latest": BREAKEVEN_MONTHLY_LATEST_FILE,
        "inflation_observed_monthly": OBSERVED_INFLATION_FILE,
        "lecap_macro_history": LECAP_MACRO_HISTORY_FILE,
        "lecap_macro_monthly": LECAP_MACRO_MONTHLY_FILE,
        "cer_curve_history": CER_CURVE_HISTORY_FILE,
        "cer_curve_latest": CER_CURVE_LATEST_FILE,
        "hard_dollar_curve_history": HARD_DOLLAR_CURVE_HISTORY_FILE,
        "hard_dollar_curve_latest": HARD_DOLLAR_CURVE_LATEST_FILE,
        "hard_dollar_cashflow_audit": HARD_DOLLAR_AUDIT_FILE,
        "dollar_linked_history": DOLLAR_LINKED_HISTORY_FILE,
        "dollar_linked_latest": DOLLAR_LINKED_LATEST_FILE,
        "monetary_variables_latest": MONETARY_VARIABLES_LATEST_FILE,
        "monetary_variables_economy_pib_history": MONETARY_VARIABLES_ECONOMY_PIB_HISTORY_FILE,
        "monetary_variables_peso_usd_pib_history": MONETARY_VARIABLES_PESO_USD_PIB_HISTORY_FILE,
        "monetary_variables_usd_stocks_history": MONETARY_VARIABLES_USD_STOCKS_HISTORY_FILE,
        "monetary_variables_block_freshness": MONETARY_VARIABLES_BLOCK_FRESHNESS_FILE,
        "monetary_variables_source_table": MONETARY_VARIABLES_SOURCE_TABLE_FILE,
        "monetary_rates_tna_series": MONETARY_RATES_TNA_SERIES_FILE,
        "monetary_rates_tna_reference": MONETARY_RATES_TNA_REFERENCE_FILE,
        "monetary_rates_tna_latest": MONETARY_RATES_TNA_LATEST_FILE,
        "fci_reported_slice_history": FCI_REPORTED_SLICE_HISTORY_FILE,
        "meta": META_FILE,
    }


if __name__ == "__main__":
    run()
