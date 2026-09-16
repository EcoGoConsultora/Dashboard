from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from src.common.cer_breakeven import (
    BREAKEVEN_HISTORY_COLUMNS,
    BREAKEVEN_MONTHLY_COLUMNS,
    build_breakeven_forward_monthly,
    build_breakeven_monitor_live,
    load_cer_index_history,
)


DEFAULT_FIXED_METRICS_PATH = Path("data/processed/lecap_monitor_history.csv")
DEFAULT_FIXED_METRICS_FALLBACK_PATH = Path("data/processed/iol_zero_coupon_metrics.csv")
DEFAULT_CER_HISTORY_PATH = Path("data/processed/cer_curve_history.csv")
DEFAULT_CER_HISTORY_FALLBACK_PATH = Path("data/raw/cer_bond_prices_history.csv")
DEFAULT_CER_INDEX_PATH = Path("data/raw/bcra_cer_index_daily.csv")
DEFAULT_BREAKEVEN_REFERENCE_PATH = Path("legacy/breakeven_reference/Infla_breakeven_2026-03-18_reference.xlsx")
PAIR_LABEL_CORRECTIONS = {
    "T17A6 - TZXM6": "S17A6 - TZXM6",
}
IGNORED_PAIR_LABELS = {"S17A6 - TZXM6"}


def build_breakeven_monitor(
    fixed_metrics_path: Path | str = DEFAULT_FIXED_METRICS_PATH,
    cer_history_path: Path | str = DEFAULT_CER_HISTORY_PATH,
    cer_index_path: Path | str = DEFAULT_CER_INDEX_PATH,
    fallback_workbook_path: Path | str = DEFAULT_BREAKEVEN_REFERENCE_PATH,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    fixed_metrics = _load_metrics_frame(
        primary_path=fixed_metrics_path,
        fallback_path=DEFAULT_FIXED_METRICS_FALLBACK_PATH,
    )
    cer_history = _load_metrics_frame(
        primary_path=cer_history_path,
        fallback_path=DEFAULT_CER_HISTORY_FALLBACK_PATH,
    )
    cer_index = load_cer_index_history(cer_index_path)
    live_history, live_latest = build_breakeven_monitor_live(
        fixed_metrics=fixed_metrics,
        cer_price_history=cer_history,
        cer_index_history=cer_index,
    )
    if not live_history.empty:
        return live_history, live_latest
    return _build_breakeven_monitor_from_workbook(path=fallback_workbook_path)


def _load_metrics_frame(primary_path: Path | str, fallback_path: Path | str) -> pd.DataFrame:
    primary = Path(primary_path)
    fallback = Path(fallback_path)
    path = primary if primary.exists() and primary.stat().st_size > 0 else fallback
    if not path.exists() or path.stat().st_size == 0:
        return pd.DataFrame()

    frame = pd.read_csv(path)
    for column in ("date", "maturity_date", "latest_cer_date", "cer_target_date", "issue_date", "payment_date"):
        if column in frame.columns:
            frame[column] = pd.to_datetime(frame[column], errors="coerce")
    return frame


def _build_breakeven_monitor_from_workbook(
    path: Path | str = DEFAULT_BREAKEVEN_REFERENCE_PATH,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    workbook_path = Path(path)
    if not workbook_path.exists() or workbook_path.stat().st_size == 0:
        empty = pd.DataFrame(columns=BREAKEVEN_HISTORY_COLUMNS)
        return empty.copy(), empty.copy()

    wb = load_workbook(workbook_path, data_only=True, read_only=False)
    try:
        sheet_name = next((name for name in wb.sheetnames if "Infl" in name), None)
        if sheet_name is None:
            raise ValueError("Breakeven reference workbook does not contain the Inflation sheet.")
        ws = wb[sheet_name]

        blocks: list[tuple[int, str, str, str]] = []
        for start_col in range(4, ws.max_column + 1, 4):
            raw_label = ws.cell(1, start_col).value
            if raw_label is None:
                continue
            pair_label = PAIR_LABEL_CORRECTIONS.get(str(raw_label).strip(), str(raw_label).strip())
            if pair_label in IGNORED_PAIR_LABELS:
                continue
            fixed_symbol, cer_symbol = _parse_pair_label(pair_label)
            blocks.append((start_col, pair_label, fixed_symbol, cer_symbol))

        rows: list[dict[str, object]] = []
        for row_idx in range(3, ws.max_row + 1):
            market_date = _as_timestamp(ws.cell(row_idx, 1).value)
            latest_cer_date = _as_timestamp(ws.cell(row_idx, 2).value)
            latest_cer_value = _as_float(ws.cell(row_idx, 3).value)
            if market_date is None:
                continue

            for start_col, pair_label, fixed_symbol, cer_symbol in blocks:
                maturity_date = _as_timestamp(ws.cell(row_idx, start_col).value)
                cer_target_date = _as_timestamp(ws.cell(row_idx, start_col + 1).value)
                implied_cer = _as_float(ws.cell(row_idx, start_col + 2).value)
                breakeven_monthly = _as_float(ws.cell(row_idx, start_col + 3).value)

                if implied_cer is None and breakeven_monthly is None:
                    continue

                horizon_days = None
                if latest_cer_date is not None and cer_target_date is not None:
                    horizon_days = int((cer_target_date - latest_cer_date).days)

                annual_equivalent = None
                if breakeven_monthly is not None:
                    annual_equivalent = (1.0 + breakeven_monthly) ** 12 - 1.0

                rows.append(
                    {
                        "date": market_date,
                        "pair_label": pair_label,
                        "fixed_symbol": fixed_symbol,
                        "cer_symbol": cer_symbol,
                        "maturity_date": maturity_date,
                        "cer_target_date": cer_target_date,
                        "latest_cer_date": latest_cer_date,
                        "latest_cer_value": latest_cer_value,
                        "implied_cer": implied_cer,
                        "breakeven_monthly": breakeven_monthly,
                        "breakeven_monthly_pct": breakeven_monthly * 100.0 if breakeven_monthly is not None else None,
                        "breakeven_annual_equivalent": annual_equivalent,
                        "horizon_days": horizon_days,
                        "is_active": bool(horizon_days is not None and horizon_days > 0),
                    }
                )
    finally:
        wb.close()

    if not rows:
        empty = pd.DataFrame(columns=BREAKEVEN_HISTORY_COLUMNS)
        return empty.copy(), empty.copy()

    history = pd.DataFrame(rows, columns=BREAKEVEN_HISTORY_COLUMNS).copy()
    history["date"] = pd.to_datetime(history["date"], errors="coerce").dt.normalize()
    history["maturity_date"] = pd.to_datetime(history["maturity_date"], errors="coerce").dt.normalize()
    history["cer_target_date"] = pd.to_datetime(history["cer_target_date"], errors="coerce").dt.normalize()
    history["latest_cer_date"] = pd.to_datetime(history["latest_cer_date"], errors="coerce").dt.normalize()
    numeric_columns = [
        "latest_cer_value",
        "implied_cer",
        "breakeven_monthly",
        "breakeven_monthly_pct",
        "breakeven_annual_equivalent",
        "horizon_days",
    ]
    for column in numeric_columns:
        history[column] = pd.to_numeric(history[column], errors="coerce")

    history = history.dropna(subset=["date", "pair_label", "cer_symbol", "breakeven_monthly"]).copy()
    history = history.sort_values(["date", "maturity_date", "pair_label"]).reset_index(drop=True)

    latest_date = history["date"].max()
    latest = (
        history[(history["date"] == latest_date) & (history["is_active"])]
        .sort_values(["maturity_date", "pair_label"])
        .reset_index(drop=True)
    )
    return history, latest


def _parse_pair_label(value: str) -> tuple[str, str]:
    parts = [part.strip().upper() for part in str(value).split("-", maxsplit=1)]
    if len(parts) != 2:
        return str(value).strip().upper(), ""
    return parts[0], parts[1]


def _as_timestamp(value: object) -> pd.Timestamp | None:
    if value is None or pd.isna(value):
        return None
    ts = pd.Timestamp(value)
    if pd.isna(ts):
        return None
    return ts.normalize()


def _as_float(value: object) -> float | None:
    if value is None or pd.isna(value):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None
