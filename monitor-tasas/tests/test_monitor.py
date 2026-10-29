import json
import unittest
from datetime import date
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

import pandas as pd

from src.breakeven.calendar import MarketCalendar
from src.breakeven.sources_rava import _parse_rava_number
from src.backfill import _merge_missing_closes, market_dates, quotes_on
from src import export_series
from src.monitor import _cer_nodes, _fixed_nodes, _history_for_html, _money_market, _publish_static_html, pending_closes


class MonitorTests(unittest.TestCase):
    def test_market_calendar_distinguishes_trading_and_settlement_closures(self):
        calendar = MarketCalendar.from_csv(Path(__file__).resolve().parents[1] / "config/calendars/ar_market_calendar_2025_2026.csv")
        self.assertFalse(calendar.is_business_day(date(2026, 11, 6)))
        self.assertEqual(market_dates(date(2026, 11, 6), date(2026, 11, 11)),
                         [date(2026, 11, 6), date(2026, 11, 11)])
        self.assertEqual(market_dates(date(2026, 12, 24), date(2026, 12, 31)),
                         [date(2026, 12, 24), date(2026, 12, 28), date(2026, 12, 29), date(2026, 12, 30)])

    def test_a3_keeps_only_peso_one_day_reference_wheels(self):
        repo = [{"fecha": "2026-09-23T00:00:00", "details": [
            {"rueda": "REPO", "moneda": "$", "plazo": "001", "tasaPP": 20.4, "volumen": 10},
            {"rueda": "REPX", "moneda": "$", "plazo": "001", "tasaPP": 40.0, "volumen": 10},
            {"rueda": "SIMU", "moneda": "$", "plazo": "001", "tasaPP": 21.0, "volumen": 5},
            {"rueda": "REPO", "moneda": "$", "plazo": "007", "tasaPP": 23.0, "volumen": 5},
        ]}]
        cauc = [{"fecha": "2026-09-23", "details": [
            {"ticker": "CAARS", "moneda": "$", "plazo": "001", "tasaPP": 20.5, "montoConcertado": 7},
            {"ticker": "CAUSD", "moneda": "D", "plazo": "001", "tasaPP": 1.0},
        ]}]
        with patch("src.monitor._get_a3", side_effect=[repo, cauc]):
            result, _ = _money_market(date(2026, 9, 23))
        self.assertEqual(set(result), {"repo_1d", "simu_1d", "caucion_1d"})
        self.assertEqual(result["repo_1d"]["tna_pct"], 20.4)
        self.assertEqual(result["simu_1d"]["tna_pct"], 21.0)

    def test_a3_overnight_on_friday_is_the_three_day_tenor(self):
        repo = [{"fecha": "2026-09-25T00:00:00", "details": [
            {"rueda": "REPO", "moneda": "$", "plazo": "003", "tasaPP": 20.44, "volumen": 10},
            {"rueda": "REPO", "moneda": "$", "plazo": "007", "tasaPP": 23.0, "volumen": 5},
        ]}]
        cauc = [{"fecha": "2026-09-25", "details": [
            {"ticker": "CAARS", "moneda": "$", "plazo": "003", "tasaPP": 20.6, "montoConcertado": 7},
        ]}]
        with patch("src.monitor._get_a3", side_effect=[repo, cauc]):
            result, _ = _money_market(date(2026, 9, 25))
        self.assertEqual(result["repo_1d"]["tna_pct"], 20.44)
        self.assertEqual(result["repo_1d"]["tenor_days"], 3)
        self.assertEqual(result["caucion_1d"]["tna_pct"], 20.6)

    def test_pending_closes_ignores_bonds_without_recent_trades(self):
        payload = {"meta": {"market_date": "2026-10-13"}, "quality": {
            "rava_errors": {"S30O6": "Cotización de 2026-10-09, se esperaba 2026-10-13",
                            "D10Y7": "Cotización de 2026-09-17, se esperaba 2026-10-13",
                            "X99Z9": "HTTPError: 503"},
            "missing_hard_dollar": ["GD30D"]}}
        self.assertEqual(pending_closes(payload), ["GD30D", "S30O6"])

    def test_fixed_cashflow_matches_audited_legacy_node(self):
        terms = [{"symbol": "S30O6", "family": "LECAP", "issue_date": "2025-10-31",
                  "maturity_date": "2026-10-30", "tem_decimal": 0.0255, "terms_source": "official"}]
        quotes = {"S30O6": {"price": 132.44, "volume_nominal": 1e9, "source_url": "rava"}}
        rows = _fixed_nodes(terms, quotes, date(2026, 9, 23), MarketCalendar(frozenset(), frozenset()))
        self.assertEqual(len(rows), 1)
        self.assertAlmostEqual(rows[0]["yield_tea_pct"], 23.98439990174599, places=8)

    def test_cer_reports_contractual_nominal_yield_only_when_fixing_known(self):
        entries = [{"symbol": "X30S6", "family": "LECER", "maturity_date": "2026-09-30",
                    "cer_base_index": 714.98, "zero_coupon": True, "terms_source": "official"}]
        quotes = {"X30S6": {"price": 117.16, "volume_nominal": 1e9, "source_url": "rava"}}
        cer = pd.DataFrame([{"date": date(2026, 9, 16), "cer_index": 842.0},
                            {"date": date(2026, 9, 24), "cer_index": 844.0}])
        rows = _cer_nodes(entries, quotes, date(2026, 9, 23), MarketCalendar(frozenset(), frozenset()), cer)
        self.assertIsNotNone(rows[0]["nominal_yield_tea_pct_if_known"])
        self.assertEqual(rows[0]["cer_target_date"], "2026-09-16")

    def test_rava_nominal_volume_uses_thousands_separator(self):
        self.assertEqual(_parse_rava_number("12.859", integer_grouping=True), 12859)
        self.assertEqual(_parse_rava_number("150.840,00"), 150840)

    def test_claude_template_becomes_self_contained_file(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            (root / "site").mkdir()
            (root / "site/template.html").write_text("<script>const data=__MONITOR_DATA_JSON__;</script>", encoding="utf-8")
            with patch("src.monitor.ROOT", root):
                self.assertTrue(_publish_static_html({"date": "2026-09-23", "text": "<x>"}))
            markup = (root / "index.html").read_text(encoding="utf-8")
            self.assertIn('"date": "2026-09-23"', markup)
            self.assertIn("\\u003cx>", markup)

    def test_history_marker_carries_only_earlier_snapshots(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            (root / "site").mkdir()
            (root / "output/history").mkdir(parents=True)
            (root / "site/template.html").write_text(
                "<script>const d=__MONITOR_DATA_JSON__;var h=__MONITOR_HISTORY_JSON__;</script>", encoding="utf-8")
            row = {"symbol": "S30N6", "maturity_date": "2026-11-30", "yield_tea_pct": 27.123456789,
                   "price_source_url": "https://x", "terms_source": "y", "known_nominal_payoff_per_vn100": None}
            for day, backfill in (("2026-09-22", True), ("2026-09-23", False), ("2026-09-24", False)):
                snapshot = {"meta": {"market_date": day, "backfill": backfill}, "curves": {"fixed": [row]},
                            "money_market": {"repo_1d": {"date": day, "tna_pct": 20.1, "volume_ars": 1}}}
                (root / f"output/history/{day}.json").write_text(json.dumps(snapshot), encoding="utf-8")
            with patch("src.monitor.ROOT", root):
                history = _history_for_html({"meta": {"market_date": "2026-09-23"}})
                self.assertTrue(_publish_static_html({"meta": {"market_date": "2026-09-23"}}))
            self.assertEqual(history["dates"], ["2026-09-22"])
            snap = history["snapshots"]["2026-09-22"]
            self.assertTrue(snap["backfill"])
            self.assertEqual(snap["curves"]["fixed"], [{"symbol": "S30N6", "maturity_date": "2026-11-30", "yield_tea_pct": 27.1235}])
            self.assertEqual(snap["money_market"]["repo_1d"], {"date": "2026-09-22", "tna_pct": 20.1})
            markup = (root / "index.html").read_text(encoding="utf-8")
            self.assertNotIn("__MONITOR_HISTORY_JSON__", markup)
            self.assertIn('"2026-09-22"', markup)

    def test_backfill_uses_only_the_close_of_the_requested_date(self):
        frame = pd.DataFrame({"date": pd.to_datetime(["2026-09-21", "2026-09-22"]), "symbol": ["S30N6", "S30N6"],
                              "price": [123.0, 124.0], "volume": [10.0, 20.0]})
        quotes, errors = quotes_on({"S30N6": frame}, {"T30J7": "HTTPError"}, {"S30N6", "T30J7"}, date(2026, 9, 22))
        self.assertEqual(quotes["S30N6"]["price"], 124.0)
        self.assertEqual(quotes["S30N6"]["volume_nominal"], 20.0)
        self.assertEqual(errors, {"T30J7": "HTTPError"})
        quotes, errors = quotes_on({"S30N6": frame}, {}, {"S30N6"}, date(2026, 9, 23))
        self.assertEqual(quotes, {})
        self.assertIn("2026-09-23", errors["S30N6"])

    def test_history_repair_preserves_original_prices_and_money_market(self):
        original = {
            "meta": {"status": "partial", "generated_at": "original"},
            "curves": {"fixed": [{"symbol": "A", "price": 101}]},
            "money_market": {"repo_1d": {"tna_pct": 25}},
            "quality": {"warnings": ["Calendario estimado", "Rava sin precio de fecha: B"],
                        "rava_errors": {"B": "cotización anterior"}, "missing_hard_dollar": [],
                        "missing_money_market": []},
        }
        rebuilt = {"meta": {"registry_version": "2026-10-08"},
                   "curves": {"fixed": [{"symbol": "A", "price": 999}, {"symbol": "B", "price": 102}]}}
        repaired, added = _merge_missing_closes(original, rebuilt)
        self.assertEqual(added, ["B"])
        self.assertEqual(repaired["curves"]["fixed"], [{"symbol": "A", "price": 101}, {"symbol": "B", "price": 102}])
        self.assertEqual(repaired["money_market"]["repo_1d"]["tna_pct"], 25)
        self.assertEqual(repaired["quality"]["rava_errors"], {})
        self.assertEqual(repaired["quality"]["warnings"], ["Calendario estimado"])
        self.assertEqual(repaired["meta"]["status"], "complete")

    def test_series_export_skips_near_maturity_yields_and_empty_wheels(self):
        def snapshot(day, cer):
            return {"meta": {"market_date": day, "backfill": True}, "money_market": {}, "quality": {},
                    "curves": {"cer": cer, "tamar": [], "dollar_linked": []}}
        near = {"symbol": "TZXM6", "maturity_date": "2026-03-31", "settlement_date": "2026-03-27",
                "real_yield_tea_pct": 40785.3, "price_ars_per_vn100": 100.0}
        far = {"symbol": "TZX27", "maturity_date": "2027-06-30", "settlement_date": "2026-03-27",
               "real_yield_tea_pct": 4.5, "price_ars_per_vn100": 200.0}
        with TemporaryDirectory() as tmp:
            history = Path(tmp)
            for day, cer in (("2026-03-26", [near, far]), ("2026-03-27", [near])):
                (history / f"{day}.json").write_text(json.dumps(snapshot(day, cer)), encoding="utf-8")
            with patch.object(export_series, "HISTORY_DIR", history):
                out = export_series.export(None, None, history / "out.xlsx")
            from openpyxl import load_workbook
            wb = load_workbook(out)
        cer = list(wb["CER"].iter_rows(min_row=6, values_only=True))
        self.assertEqual(len([r for r in cer if r[0] is not None and hasattr(r[0], "year")]), 1)
        self.assertIsNone(cer[0][1])  # TZXM6 a cuatro días del vencimiento
        self.assertAlmostEqual(cer[0][2], 0.045)
        data = list(wb["Datos"].iter_rows(min_row=5, values_only=True))
        self.assertEqual(len(data), 3)
        self.assertEqual([r[-1] for r in data if r[2] == "TZXM6"], ["No", "No"])


if __name__ == "__main__":
    unittest.main()
