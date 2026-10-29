"""Exporta a Excel las series diarias de rendimientos CER, TAMAR y dólar linked.

Lee los snapshots de `output/history` (una rueda por archivo) y no recalcula nada: cada
rendimiento es el que el monitor publicó o reconstruyó para esa fecha de mercado.

    py -3.14 -m src.export_series
    py -3.14 -m src.export_series --start 2026-06-01 --end 2026-10-15 --out serie.xlsx
"""

from __future__ import annotations

import argparse
import json
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
HISTORY_DIR = ROOT / "output/history"
THEME_PATH = ROOT / "config/brand/ecogo-theme1.xml"

# Paleta y tipografías Eco Go (mismas que la skill ecogo-deck).
TEAL, TEAL_DARK, TEAL_SOFT, INK, MUTED, ROW_ALT = "2890A4", "1B5F5E", "DCEEEE", "333333", "7A8C99", "F4F7F8"
FONT_HEAD, FONT_BODY = "HK Grotesk SemiBold", "HK Grotesk"
PCT, DATE_FMT = "0.00%", "dd/mm/yyyy"
# Con pocos días al vencimiento la TIR anualizada se dispara (TZXM6 llegó a 40.785% a un día
# de vencer): en las hojas de series se omite. La hoja Datos conserva todas las observaciones.
MIN_DAYS = 30

GROUPS = {
    "cer": {
        "sheet": "CER",
        "title": "Bonos CER: TIR real",
        "unit": "TIR real TEA, en %, por especie. Columnas ordenadas por vencimiento",
        "yield": "real_yield_tea_pct",
        "price": "price_ars_per_vn100",
    },
    "tamar": {
        "sheet": "TAMAR",
        "title": "Bonos TAMAR: TIR",
        "unit": "TIR TEA, en %, por especie, y TAMAR bancos privados BCRA (TNA). Columnas ordenadas por vencimiento",
        "yield": "tir_tea_pct",
        "price": "price_per_vn_100",
    },
    "dollar_linked": {
        "sheet": "Dollar linked",
        "title": "Bonos dólar linked: rendimiento en USD",
        "unit": "TEA en USD equivalentes, en %, por especie, y A3500 BCRA. Columnas ordenadas por vencimiento",
        "yield": "usd_equivalent_yield_tea_pct",
        "price": "price_ars_per_vn100_usd",
    },
}
SOURCE = "Eco Go en base a Rava, BCRA (CER, TAMAR, A3500) y Secretaría de Finanzas"


def load_snapshots(start: date | None, end: date | None) -> list[dict]:
    snapshots = []
    for path in sorted(HISTORY_DIR.glob("*.json")):
        day = date.fromisoformat(path.stem)
        if (start and day < start) or (end and day > end):
            continue
        snapshots.append(json.loads(path.read_text(encoding="utf-8")))
    if not snapshots:
        raise SystemExit(f"No hay snapshots en {HISTORY_DIR} para el rango pedido")
    return snapshots


def _pct(value) -> float | None:
    return None if value is None else value / 100


def _font(bold=False, size=10.5, color=INK) -> Font:
    return Font(name=FONT_HEAD if bold else FONT_BODY, bold=bold, size=size, color=color)


def _new_sheet(wb: Workbook, name: str, periodo: str):
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = TEAL
    ws.page_setup.orientation, ws.page_setup.paperSize = "landscape", ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    # Pie teal en negrita sin nombrar la fuente: con HK Grotesk Excel rasteriza el pie.
    for footer in (ws.oddFooter, ws.evenFooter):
        for part, text in ((footer.left, f"ECOGO S.A.  ·  {periodo}"), (footer.center, f"Fuente: {SOURCE}"), (footer.right, "Pág. &P de &N")):
            part.text, part.size, part.color = "&B" + text, 9, TEAL
    return ws


def _title(ws, title: str, unit: str) -> None:
    ws["A1"], ws["A2"] = title, unit.upper()
    ws["A1"].font = Font(name=FONT_HEAD, bold=True, size=13, color="000000")
    ws["A2"].font = _font(size=9, color=MUTED)


def _source(ws, row: int, text: str = SOURCE) -> None:
    cell = ws.cell(row, 1, f"Fuente: {text}")
    cell.font = Font(name=FONT_BODY, size=9, color=TEAL)


def _header(ws, row: int, values: list, first_left=True) -> None:
    for col, value in enumerate(values, start=1):
        cell = ws.cell(row, col, value)
        cell.font = _font(bold=True, size=10.5, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=TEAL_DARK)
        cell.alignment = Alignment(horizontal="left" if col == 1 and first_left else "right", vertical="center", wrap_text=True)


def _data_cell(ws, row: int, col: int, value, fmt: str | None, zebra: bool) -> None:
    cell = ws.cell(row, col, value)
    cell.font = _font()
    cell.alignment = Alignment(horizontal="left" if col == 1 else "right", vertical="center")
    if fmt:
        cell.number_format = fmt
    if zebra:
        cell.fill = PatternFill("solid", fgColor=ROW_ALT)


def _days_to_maturity(row: dict, market_date: str) -> int:
    settlement = row.get("settlement_date") or market_date
    return (date.fromisoformat(row["maturity_date"]) - date.fromisoformat(settlement)).days


def _fmt(iso: str) -> str:
    return date.fromisoformat(iso).strftime("%d/%m/%Y")


def _instruments(snapshots: list[dict], group: str) -> list[tuple[str, str]]:
    maturities = {}
    for snap in snapshots:
        for row in snap["curves"].get(group, []):
            maturities[row["symbol"]] = row["maturity_date"]
    return sorted(maturities.items(), key=lambda item: (item[1], item[0]))


def _wide_sheet(wb: Workbook, snapshots: list[dict], group: str, periodo: str, min_days: int) -> int:
    spec = GROUPS[group]
    ws = _new_sheet(wb, spec["sheet"], periodo)
    _title(ws, spec["title"], spec["unit"])
    instruments = _instruments(snapshots, group)

    extra: list[tuple[str, str, object]] = []  # (encabezado, formato, función snapshot -> valor)
    if group == "tamar":
        tamar = lambda s: (s.get("money_market") or {}).get("tamar_private") or {}
        extra = [("TAMAR BCRA\n(TNA)", PCT, lambda s: _pct(tamar(s).get("tna_pct"))),
                 ("Fecha dato\nTAMAR", DATE_FMT, lambda s: date.fromisoformat(tamar(s)["date"]) if tamar(s).get("date") else None)]
    if group == "dollar_linked":
        def a3500(s):
            rows = s["curves"].get("dollar_linked") or []
            return rows[0].get("fx_a3500_ars_per_usd") if rows else None
        extra = [("A3500\n($/USD)", "#,##0.00", a3500)]

    head_row = 4
    _header(ws, head_row, ["Fecha"] + [h for h, _, _ in extra] + [symbol for symbol, _ in instruments])
    _header(ws, head_row + 1, ["Vencimiento"] + ["" for _ in extra] + [date.fromisoformat(m) for _, m in instruments])
    for col in range(2 + len(extra), 2 + len(extra) + len(instruments)):
        ws.cell(head_row + 1, col).number_format = "dd/mm/yy"
    ws.row_dimensions[head_row].height = 30

    first_col = 2 + len(extra)
    row = head_row + 2
    for snap in snapshots:
        market_date = snap["meta"]["market_date"]
        yields = {r["symbol"]: r.get(spec["yield"]) for r in snap["curves"].get(group, [])
                  if _days_to_maturity(r, market_date) >= min_days}
        extras = [getter(snap) for _, _, getter in extra]
        values = [_pct(yields.get(symbol)) for symbol, _ in instruments]
        # Rueda sin cierres del grupo (o anterior a la primera especie). En TAMAR se conserva
        # si hay fijación BCRA, que es una serie propia.
        if all(v is None for v in values) and not (group == "tamar" and extras[0] is not None):
            continue
        zebra = (row - head_row) % 2 == 1
        _data_cell(ws, row, 1, date.fromisoformat(market_date), DATE_FMT, zebra)
        for j, ((_, fmt, _), value) in enumerate(zip(extra, extras)):
            _data_cell(ws, row, 2 + j, value, fmt, zebra)
        for j, value in enumerate(values):
            _data_cell(ws, row, first_col + j, value, PCT, zebra)
        row += 1

    last_row = row - 1
    ws.column_dimensions["A"].width = 13
    for col in range(2, first_col + len(instruments)):
        ws.column_dimensions[get_column_letter(col)].width = 11.5
    ws.freeze_panes = ws.cell(head_row + 2, 2)
    ws.print_title_rows = f"{head_row}:{head_row + 1}"
    ws.auto_filter.ref = f"A{head_row + 1}:{get_column_letter(first_col + len(instruments) - 1)}{last_row}"
    _source(ws, last_row + 2)
    return len(instruments)


DATA_COLUMNS = [
    ("Fecha", 12, DATE_FMT), ("Familia", 14, None), ("Especie", 10, None), ("Tipo", 12, None),
    ("Vencimiento", 12, DATE_FMT), ("Días al vto.", 10, "#,##0"), ("Precio (VN 100)", 13, "#,##0.00"),
    ("Rendimiento TEA", 13, PCT), ("Medida", 26, None), ("TIR TNA (TAMAR)", 13, PCT),
    ("TIR TEM", 10, PCT), ("Duración modif. (años)", 12, "0.00"), ("Volumen nominal", 17, "#,##0"),
    ("Rueda reconstruida", 12, None), ("En hojas de series", 12, None),
]
MEASURE = {"cer": "TIR real TEA (CER)", "tamar": "TIR TEA (proyección TAMAR)", "dollar_linked": "TEA en USD equivalentes"}
FAMILY = {"cer": "CER", "tamar": "TAMAR", "dollar_linked": "Dólar linked"}


def _data_sheet(wb: Workbook, snapshots: list[dict], periodo: str, min_days: int) -> int:
    ws = _new_sheet(wb, "Datos", periodo)
    _title(ws, "Base de datos: rendimientos por especie y rueda", "Formato largo, una fila por especie y fecha de mercado")
    _header(ws, 4, [name for name, _, _ in DATA_COLUMNS])
    ws.row_dimensions[4].height = 30
    row = 5
    for snap in snapshots:
        market_date = date.fromisoformat(snap["meta"]["market_date"])
        rebuilt = "Sí" if snap["meta"].get("backfill") else "No"
        for group, spec in GROUPS.items():
            for r in snap["curves"].get(group, []):
                maturity = date.fromisoformat(r["maturity_date"])
                days = _days_to_maturity(r, market_date.isoformat())
                tem = r.get("tir_tem_pct") if group == "tamar" else r.get("yield_tem_pct")
                values = [market_date, FAMILY[group], r["symbol"], r.get("family") or r.get("instrument_type"),
                          maturity, days, r.get(spec["price"]), _pct(r.get(spec["yield"])),
                          MEASURE[group], _pct(r.get("tir_tna_pct")), _pct(tem), r.get("modified_duration_years"),
                          r.get("volume_nominal"), rebuilt, "Sí" if days >= min_days else "No"]
                for col, (value, (_, _, fmt)) in enumerate(zip(values, DATA_COLUMNS), start=1):
                    _data_cell(ws, row, col, value, fmt, False)
                row += 1
    for col, (_, width, _) in enumerate(DATA_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A5"
    ws.print_title_rows = "4:4"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(DATA_COLUMNS))}{row - 1}"
    return row - 5


def _notes_sheet(wb: Workbook, snapshots: list[dict], periodo: str, counts: dict, min_days: int) -> None:
    ws = _new_sheet(wb, "Notas", periodo)
    rebuilt = [s["meta"]["market_date"] for s in snapshots if s["meta"].get("backfill")]
    daily = [s["meta"]["market_date"] for s in snapshots if not s["meta"].get("backfill")]
    _title(ws, "Series de rendimientos CER, TAMAR y dólar linked", f"Ruedas del {periodo}")
    lines = [
        ("Contenido", None),
        ("CER", f"TIR real TEA de cada LECER/BONCER cero cupón ({counts['cer']} especies). Aproximada con el CER de liquidación; no incluye bonos CER con cupones."),
        ("TAMAR", f"TIR TEA de cada bono/letra TAMAR puro ({counts['tamar']} especies). Condicional a la proyección de fijaciones futuras (media simple de las últimas cinco TAMAR publicadas). Incluye la TAMAR bancos privados BCRA (TNA) con su fecha, que puede rezagarse uno o dos días."),
        ("Dollar linked", f"Rendimiento TEA en USD equivalentes de cada dólar linked cero cupón ({counts['dollar_linked']} especies), con la A3500 vigente. No proyecta el tipo de cambio: un valor negativo indica que el bono paga menos que el dólar oficial actual."),
        ("Datos", f"Las {counts['rows']:,} observaciones en formato largo".replace(",", ".") + ", con precio, días al vencimiento, duración y volumen."),
        ("", None),
        ("Metodología", None),
        ("Precios", "Cierre de Rava de cada fecha de mercado. Una especie sin cierre de esa fecha queda vacía: no se arrastra un precio anterior."),
        ("Plazo mínimo", f"En las hojas CER, TAMAR y Dollar linked se omite cada especie en las ruedas con menos de {min_days} días al vencimiento: anualizar un plazo tan corto da tasas sin sentido económico. La hoja Datos las conserva (columna \"En hojas de series\")."),
        ("Especies", "Sólo especies con condiciones de emisión verificadas contra Secretaría de Finanzas / Boletín Oficial. Cada columna arranca con la primera cotización registrada y termina al vencimiento."),
        ("Ruedas", f"{len(snapshots)} ruedas. {len(rebuilt)} reconstruidas con cierres históricos ({_fmt(rebuilt[0]) if rebuilt else '-'} a {_fmt(rebuilt[-1]) if rebuilt else '-'}) y {len(daily)} de la corrida diaria ({_fmt(daily[0]) if daily else '-'} a {_fmt(daily[-1]) if daily else '-'}). La columna \"Rueda reconstruida\" de Datos lo indica."),
        ("Unidades", "Tasas en % efectivo anual (TEA) salvo que se indique TNA o TEM."),
        ("", None),
        ("Fuente", SOURCE),
    ]
    row = 4
    for label, text in lines:
        if text is None and label:
            _header(ws, row, [label, ""])
        elif label:
            ws.cell(row, 1, label).font = _font(bold=True, color=TEAL_DARK)
            ws.cell(row, 1).fill = PatternFill("solid", fgColor=TEAL_SOFT)
            cell = ws.cell(row, 2, text)
            cell.font = _font()
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            ws.cell(row, 1).alignment = Alignment(vertical="top")
            ws.row_dimensions[row].height = 15 * max(1, -(-len(text) // 95))
        row += 1
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 100


def export(start: date | None, end: date | None, out: Path | None, min_days: int = MIN_DAYS) -> Path:
    snapshots = load_snapshots(start, end)
    first, last = snapshots[0]["meta"]["market_date"], snapshots[-1]["meta"]["market_date"]
    periodo = f"{_fmt(first)} al {_fmt(last)}"
    out = out or ROOT / "output/exports" / f"Series_CER_TAMAR_DL_{last}.xlsx"

    wb = Workbook()
    wb.remove(wb.active)
    if THEME_PATH.exists():
        wb.loaded_theme = THEME_PATH.read_bytes()  # tema oficial Ecogo: paleta y HK Grotesk
    counts = {group: _wide_sheet(wb, snapshots, group, periodo, min_days) for group in GROUPS}
    counts["rows"] = _data_sheet(wb, snapshots, periodo, min_days)
    _notes_sheet(wb, snapshots, periodo, counts, min_days)  # usa los conteos, pero va primera
    wb.move_sheet("Notas", offset=-len(wb.sheetnames) + 1)
    wb.active = 0
    wb.properties.creator = "Eco Go · ECOGO S.A."
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def main() -> None:
    parser = argparse.ArgumentParser(description="Exporta a Excel las series CER, TAMAR y dólar linked del monitor")
    parser.add_argument("--start", type=date.fromisoformat)
    parser.add_argument("--end", type=date.fromisoformat)
    parser.add_argument("--out", type=Path)
    parser.add_argument("--min-days", type=int, default=MIN_DAYS, help="Días mínimos al vencimiento para las hojas de series")
    args = parser.parse_args()
    print(export(args.start, args.end, args.out, args.min_days))


if __name__ == "__main__":
    main()
