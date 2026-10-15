"""Estilo EcoGo para los Excel generados con openpyxl (copia de procesos/xlsx_marca.py)."""
from __future__ import annotations

from openpyxl.styles import Alignment, Font, PatternFill

TEAL = "3C9794"
TEAL_DARK = "1B5F5E"
TEAL_LIGHT = "8FCCCA"
GRAY_DARK = "333333"
GRAY_LIGHT = "F1F1F1"
FONT = "HK Grotesk"


def celda(ws, fila: int, col: int, valor, *, bold=False, italic=False, size=10, color=GRAY_DARK,
          fill=None, fmt=None, h="left", wrap=False, v=None):
    c = ws.cell(row=fila, column=col, value=valor)
    c.font = Font(name=FONT, bold=bold, italic=italic, size=size, color=color)
    if fill:
        c.fill = PatternFill("solid", fgColor=fill)
    if fmt:
        c.number_format = fmt
    c.alignment = Alignment(horizontal=h, vertical=v or "center", wrap_text=wrap)
    return c


def encabezados(ws, fila: int, textos: list[str]) -> None:
    for j, t in enumerate(textos, start=1):
        celda(ws, fila, j, t, bold=True, size=9, color="FFFFFF", fill=TEAL_DARK, h="center", wrap=True)
    ws.row_dimensions[fila].height = 30


def titulo(ws, texto: str, sub: str, ancho: int) -> None:
    celda(ws, 1, 1, texto, bold=True, size=14, color=TEAL)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ancho)
    celda(ws, 2, 1, sub, italic=True)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ancho)


def pie(ws, fila: int, fuente: str, ancho: int, nota: str | None = None) -> None:
    if nota:
        celda(ws, fila, 1, nota, italic=True, size=9)
        ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=ancho)
        fila += 1
    fila += 1
    celda(ws, fila, 1, fuente, italic=True, size=9)
    ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=ancho)
    celda(ws, fila + 1, 1, "www.ecogo.com.ar", italic=True, size=9, color=TEAL)
